from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "sax-audit-responses-2026-10-06"
DOWNLOADS = Path("/Users/Apple/Downloads")

SOURCE_JE = DOWNLOADS / "Journal Entries Internal Control Questions for Client.docx"
SOURCE_IC = DOWNLOADS / "Internal Control Questions for Client.docx"
SOURCE_ITGC = DOWNLOADS / "IT General Controls - Questions for Client.docx"

OUT_JE = OUTPUT_DIR / "SAX Journal Entries Internal Controls Draft Responses.docx"
OUT_IC = OUTPUT_DIR / "SAX Financial Process Internal Controls Draft Responses.docx"
OUT_ITGC = OUTPUT_DIR / "SAX IT General Controls Draft Responses and Remediation Plan.docx"

NAVY = "17365D"
BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "666666"
GREEN = "237A57"
AMBER = "9C6500"
RED = "A61B1B"
BLACK = "000000"
WHITE = "FFFFFF"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), BORDER)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row) -> None:
    repeat_table_header(row)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Arial", size=10.5, color=BLACK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=5, line=1.08) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def configure_document(doc: Document, title: str, subtitle: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, before, after in (
        ("Title", 24, 0, 10),
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 13, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_p = doc.add_paragraph(style="Title")
    title_run = title_p.add_run(title)
    set_run_font(title_run, size=24, bold=True)
    subtitle_p = doc.add_paragraph()
    subtitle_run = subtitle_p.add_run(subtitle)
    set_run_font(subtitle_run, size=11.5, color=MID_GRAY)
    set_paragraph_spacing(subtitle_p, after=14, line=1.15)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("urbanXtracts  |  SAX audit response draft")
    set_run_font(hr, size=8.5, color=MID_GRAY, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Management review required before submission  |  6 October 2026  |  Page ")
    set_run_font(fr, size=8.5, color=MID_GRAY)
    add_page_field(fp)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8.5, color=MID_GRAY)


def add_intro(doc: Document, scope: str, conclusion: str) -> None:
    p = doc.add_paragraph()
    lead = p.add_run("Purpose. ")
    set_run_font(lead, bold=True)
    text = p.add_run(scope)
    set_run_font(text)
    set_paragraph_spacing(p, after=7, line=1.13)

    p = doc.add_paragraph()
    lead = p.add_run("Current conclusion. ")
    set_run_font(lead, bold=True)
    text = p.add_run(conclusion)
    set_run_font(text)
    set_paragraph_spacing(p, after=9, line=1.13)

    p = doc.add_paragraph()
    lead = p.add_run("Submission status. ")
    set_run_font(lead, bold=True, color=RED)
    text = p.add_run(
        "This is a management-review draft, not a completed management representation. "
        "Items marked Confirmation required or Control gap must be resolved or expressly disclosed before the response is sent to SAX."
    )
    set_run_font(text)
    set_paragraph_spacing(p, after=12, line=1.13)


def add_bullets(doc: Document, items: Iterable[str], level=0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        r = p.add_run(item)
        set_run_font(r)
        set_paragraph_spacing(p, after=3, line=1.08)


def add_summary_table(doc: Document, headers: List[str], rows: List[List[str]], widths: List[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size=9.5, color=WHITE, bold=True)
        set_paragraph_spacing(p, after=0, line=1.0)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, text in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            set_cell_shading(cell, WHITE if row_idx % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_run_font(run, size=9.3, bold=(i == 0))
            set_paragraph_spacing(p, after=0, line=1.05)
    doc.add_paragraph()


def status_color(status: str) -> str:
    lower = status.lower()
    if "supported" in lower and "partial" not in lower:
        return GREEN
    if "gap" in lower or "not implemented" in lower:
        return RED
    return AMBER


def add_qa(doc: Document, number: str, question: str, answer: str, status: str, evidence: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    nr = p.add_run(f"{number}  ")
    set_run_font(nr, size=10.5, color=BLUE, bold=True)
    qr = p.add_run(question)
    set_run_font(qr, size=10.5, bold=True)

    p = doc.add_paragraph()
    lead = p.add_run("Draft response. ")
    set_run_font(lead, bold=True)
    ar = p.add_run(answer)
    set_run_font(ar)
    set_paragraph_spacing(p, after=3, line=1.12)

    p = doc.add_paragraph()
    lead = p.add_run("Status. ")
    set_run_font(lead, bold=True)
    sr = p.add_run(status)
    set_run_font(sr, color=status_color(status), bold=True)
    set_paragraph_spacing(p, after=2, line=1.05)

    if evidence:
        p = doc.add_paragraph()
        lead = p.add_run("Evidence requested. ")
        set_run_font(lead, bold=True)
        er = p.add_run(evidence)
        set_run_font(er, color=MID_GRAY)
        set_paragraph_spacing(p, after=4, line=1.08)


def source_questions(path: Path) -> List[Tuple[int, str]]:
    doc = Document(path)
    result = []
    for i, paragraph in enumerate(doc.paragraphs):
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if "?" in text or text.lower().startswith("document ") or text.lower().startswith("list the "):
            result.append((i, text.lstrip("• ").strip()))
    return result


def build_journal_entries() -> None:
    doc = Document()
    configure_document(
        doc,
        "Journal Entries Internal Controls Draft Responses",
        "Prepared for management review in response to the SAX questionnaire",
    )
    add_intro(
        doc,
        "This document provides proposed responses to the journal-entry questionnaire. It reflects the known system boundary: QuickBooks Online is the accounting system of record, while UX OS displays selected financial information and does not create, post, approve, or modify QuickBooks journal entries.",
        "The portal does not introduce an automated journal-entry risk. The remaining journal-entry control assertions depend on Finance's actual QuickBooks roles, approval practice, close process, and review evidence. Those facts were not contained in the supplied questionnaire or portal repository and must be confirmed by the Finance/Controller owner.",
    )

    doc.add_heading("Control assessment summary", level=1)
    add_summary_table(
        doc,
        ["Area", "Draft assessment", "Required evidence"],
        [
            ["Portal boundary", "Supported: UX OS has no QuickBooks write route and cannot post journal entries.", "Connector configuration and current financial integration documentation."],
            ["Preparation and posting", "Confirmation required: authorized QuickBooks roles and actual preparers were not provided.", "QuickBooks user/role listing and journal report."],
            ["Approval and review", "Control gap unless evidenced: an independent documented approval and monthly review process is not presently documented.", "Journal approval evidence, close checklist, and Audit Log review."],
            ["Period-end accuracy", "Confirmation required: reconciliations, cutoff review, and recurring/non-standard entry review must be evidenced by Finance.", "Reconciliations, support, and period-close sign-off."],
        ],
        [1.35, 3.05, 2.45],
    )

    answers: Dict[int, Tuple[str, str, str]] = {
        0: (
            "Based on information provided to date, no inappropriate or unusual journal-entry activity has been reported. Finance should not finalize this response until the Controller reviews the period's QuickBooks Journal report and Audit Log for unusual users, dates, accounts, descriptions, round-dollar amounts, late postings, and entries made or changed by privileged users.",
            "Confirmation required",
            "Period Journal report, QuickBooks Audit Log, list of unusual-entry criteria, and Controller sign-off.",
        ),
        2: (
            "Proposed process: authorized Finance personnel prepare standard and non-standard journal entries in QuickBooks Online using supporting documentation. Non-standard, manual, period-end, consolidation, correction, and management entries are reviewed and approved by the Controller or another authorized reviewer who did not prepare the entry. The reviewer verifies business purpose, account coding, amount, period, and support before or promptly after posting. Recurring entries are supported by approved schedules and are revalidated at least annually. UX OS does not create or post journal entries.",
            "Confirmation required",
            "Written journal-entry policy, role matrix, sample standard and non-standard entries, and evidence of review.",
        ),
        6: (
            "Finance initiates a journal entry when a reconciliation, close schedule, correction, accrual, allocation, or approved management estimate identifies the need. The preparer assembles the calculation and source support and records the proposed debit and credit. The Controller or designated approver authorizes entries according to an approved authority matrix. The actual preparer and approver names and thresholds must be confirmed.",
            "Confirmation required",
            "Current Finance organization chart, approval matrix, and sample entry support.",
        ),
        8: (
            "Journal entries are recorded directly in QuickBooks Online by an individually authorized Finance user or generated by a documented QuickBooks transaction process. Each manual entry should include the date, accounts, debit and credit amounts, description, preparer, approver, and supporting schedule. Posting outside an open accounting period should be restricted or separately approved.",
            "Confirmation required",
            "QuickBooks configuration, user list, sample Journal report, and closed-period settings.",
        ),
        10: (
            "QuickBooks Online posts approved entries to the general ledger. Finance reviews the Journal report and account reconciliations during the monthly close. The Controller investigates unexpected accounts, periods, amounts, or users and documents resolution. UX OS is read-only for selected QuickBooks records and does not alter the ledger.",
            "Partially supported - Finance confirmation required",
            "Monthly close checklist, Journal report review, account reconciliations, and QuickBooks Audit Log.",
        ),
        15: (
            "Recommended policy: only specifically named Finance users may prepare journal entries, and only the Controller or a formally delegated approver may authorize non-standard, period-end, correcting, or management entries. Materiality-based thresholds and prohibited self-approval should be documented. Administrator access should not remove the requirement for independent review.",
            "Control design required",
            "Approved journal-entry policy and authority matrix.",
        ),
        16: (
            "Each non-standard journal entry should be approved through retained evidence before posting when the system supports it, or reviewed and signed off promptly after posting where QuickBooks workflow limitations apply. The reviewer must be independent of the preparer and verify support, account, amount, date, period, and business purpose. Evidence may be an approved close checklist, attached support, or a linked approval record.",
            "Control design required",
            "Samples showing preparer, approver, date, support, and review conclusion.",
        ),
        18: (
            "Accuracy is supported by required source documentation, balanced debits and credits, account-reconciliation review, comparison to approved schedules, and an independent review of non-standard entries. Spreadsheet calculations should be checked and versioned. Entries identified through reconciliation are corrected through a separately supported entry rather than overwriting evidence.",
            "Confirmation required",
            "Sample support, reconciliations, and review evidence.",
        ),
        19: (
            "The authorization policy should match the response above: named Finance preparers, Controller approval for non-standard and period-end entries, documented thresholds, no self-approval, and independent review of any entry posted by a privileged user. Actual limits and delegates must be approved by management.",
            "Control design required",
            "Approved authority matrix and QuickBooks role listing.",
        ),
        21: (
            "During monthly close, Finance reviews the Journal report, account reconciliations, cutoff support, and open-period settings. The reviewer verifies that each entry agrees to its support, uses the correct account and class or other approved dimension, and falls in the proper reporting period. Exceptions are corrected before the close is approved.",
            "Confirmation required",
            "Monthly close checklist, account reconciliations, cutoff review, and locked-period settings.",
        ),
        22: (
            "Authorization and accuracy should be evidenced by unique QuickBooks user IDs, role-based access, a retained preparer/approver record, supporting documents, monthly Journal report review, and periodic review of the QuickBooks Audit Log and user list. The UX OS portal adds no posting route, so all ledger authorization remains inside Finance and QuickBooks.",
            "Partially supported - Finance confirmation required",
            "QuickBooks user access report, Audit Log, entry samples, and review sign-offs.",
        ),
    }

    doc.add_heading("Questionnaire responses", level=1)
    for position, (idx, question) in enumerate(source_questions(SOURCE_JE), 1):
        answer, status, evidence = answers[idx]
        add_qa(doc, f"JE-{position:02d}", question, answer, status, evidence)

    doc.add_page_break()
    doc.add_heading("Management confirmations before submission", level=1)
    add_bullets(
        doc,
        [
            "Name the QuickBooks journal-entry preparers, approvers, and administrators, including titles.",
            "Approve entry-type and dollar thresholds and prohibit self-approval for non-standard entries.",
            "Confirm whether QuickBooks periods are closed or password-protected after the monthly close.",
            "Confirm how recurring entries, spreadsheets, and management estimates are reviewed.",
            "Review the period Journal report and Audit Log and document the conclusion on unusual activity.",
            "Provide at least one standard and one non-standard entry sample with support and approval evidence.",
        ],
    )
    doc.add_heading("Recommended immediate remediation", level=1)
    add_bullets(
        doc,
        [
            "Adopt a one-page journal-entry policy and approval matrix within 30 days.",
            "Add a journal-entry review step to the monthly close checklist and retain the reviewed report.",
            "Perform a quarterly QuickBooks user and privileged-access review.",
            "Require independent review of every entry posted by an administrator or outside normal close dates.",
        ],
    )
    doc.save(OUT_JE)


ITGC_ANSWERS: Dict[int, Tuple[str, str, str]] = {
    0: (
        "IT risk is presently assessed through implementation reviews, release-readiness checks, security tests, vendor configuration work, and management decisions made as systems change. The process is project-based rather than a formally approved enterprise IT risk assessment. Management should establish an annual assessment and require an update after any material financial-system, integration, identity, or hosting change.",
        "Partial control - formalization required",
        "Approved IT risk assessment, risk register, and annual management sign-off.",
    ),
    1: (
        "The technical owner currently coordinates UX OS architecture, integrations, access configuration, and releases. Executive management sets business priorities. Finance/Controller must approve financial-reporting rules and QuickBooks-related decisions, while Operations owns Canix and operational workflow decisions. Management should confirm the named individuals and any outsourced providers.",
        "Confirmation required",
        "Organization chart, system-owner register, and administrator listings.",
    ),
    2: (
        "There is no evidenced recurring company-wide IT risk assessment cadence. Recommended cadence is annually, with interim reassessment after a material system implementation, major integration change, security incident, or significant vendor change. High-risk remediation should be reviewed quarterly until closed.",
        "Control gap",
        "Annual control calendar and completed first assessment.",
    ),
    3: (
        "Current identification methods include system design review, release-readiness testing, server-side authorization tests, integration failure handling, deferred-item tracking, and issue reports. The formal process should score confidentiality, integrity, availability, financial-reporting impact, likelihood, vendor dependence, and recovery requirements and assign an owner and due date to each risk.",
        "Partial control - formalization required",
        "Risk methodology, scored risk register, and remediation tracking.",
    ),
    4: (
        "The entity relies on cloud and SaaS providers, including QuickBooks, Google Workspace, Supabase, Canix, GitHub, portal hosting, document scanning, email delivery, and historically Monday.com. Technical safeguards exist for several integrations, but a complete annual vendor-control review is not evidenced. Management should maintain a vendor register, review contracts and relevant SOC reports, document subprocessor and breach-notification terms, and track service availability and incidents.",
        "Partial control - vendor program required",
        "Vendor register, contracts, SOC reports or security attestations, and annual review sign-off.",
    ),
    5: (
        "Business and security priorities are set by management and implemented by the technical owner. Work is generally prioritized by production impact, security exposure, compliance or financial-reporting effect, and operational urgency. A formal intake and approval queue with documented priority, owner, due date, and closure evidence should replace informal prioritization.",
        "Partial control - formalization required",
        "Approved prioritization criteria and current IT work register.",
    ),
    6: (
        "The portal has release-readiness checks, automated security contracts, audit/event records, and integration-health views. There is no evidenced annual management review covering all in-scope IT general controls. Recommended practice is quarterly review of access, changes, incidents, backups, vulnerabilities, and vendors, plus an annual owner certification.",
        "Control gap",
        "Quarterly IT control review package and annual certification.",
    ),
    7: (
        "Portal users can report order and portal issues, and administrators have work queues for operational items. A company-wide service desk with ticket numbering, severity, assignment, service targets, root-cause classification, and closure evidence is not evidenced. Management should designate one controlled intake channel and retain the ticket history.",
        "Partial control - service process required",
        "Help-desk procedure, ticket export, service targets, and sample closed tickets.",
    ),
    8: (
        "A significant ongoing change is the UX OS portal, which consolidates controlled views and workflows while retaining QuickBooks as the accounting authority and Canix as the physical/compliance inventory authority. Portal-native workflows are replacing selected Monday workflows, while Monday records are retained and not deleted. Store onboarding, cost-object/lot controls, and some financial definitions remain on hold. Management must confirm other planned QuickBooks, payroll, network, endpoint, or staffing changes.",
        "Supported for UX OS - broader confirmation required",
        "Approved roadmap, cutover plan, system inventory, and management confirmation of other changes.",
    ),
    9: (
        "Yes. UX OS has been newly implemented and materially modified during the period. It reads selected QuickBooks data but does not create invoices, payments, journal entries, or other accounting records. It reads Canix inventory and maintains portal-owned workflows and projections. Management should document the implementation date, production scope, testing, approvals, and whether the system is financially relevant for the audit period.",
        "Supported - dates and scope require confirmation",
        "Implementation approval, release evidence, test results, and current production architecture.",
    ),
    16: (
        "Changes are initiated from business requests, identified defects, meeting decisions, release-readiness findings, and security requirements. For the portal, requirements are implemented in a Git repository and tested before release. A single mandatory change ticket containing business purpose, risk, approval, test evidence, rollback steps, and implementation result is not yet evidenced for every change.",
        "Partial control - formalization required",
        "Change tickets, issue register, and linked commits/releases.",
    ),
    17: (
        "The technical owner implements portal and integration changes. Business owners from Finance, Operations, Sales, Quality, or Administration provide requirements and should approve changes affecting their area. Exact names, titles, and delegates must be confirmed in a system-owner and change-approval matrix.",
        "Confirmation required",
        "RACI matrix and named system owners.",
    ),
    20: (
        "Recommended control: the relevant business owner authorizes each change before production, and Finance/Controller separately approves changes affecting financial data, mappings, or reports. Emergency changes require documented retrospective approval within one business day. Current evidence does not show that every production change has this approval record.",
        "Control gap",
        "Approved change policy and samples with authorization.",
    ),
    21: (
        "The portal uses source control, build verification, release-contract tests, service tests, security tests, and deployment evidence. To prove that the implemented change matches the approval, each release should link the approved request to the exact commit, test result, production version, and post-deployment validation. This linkage is not yet consistently evidenced as a formal control.",
        "Partial control - linkage required",
        "Change-to-commit trace, test result, release record, and post-release check.",
    ),
    22: (
        "Production and development access should be restricted to named users through GitHub, Supabase, hosting, and vendor administrator roles. Secrets are held server-side rather than in browser code. The current repository does not contain a complete approved access listing or evidence of separation between developer and production deployer. Management must export current privileged users and define who may approve and deploy.",
        "Partial control - access evidence required",
        "GitHub, Supabase, hosting, Google, QuickBooks, and Canix privileged-user exports.",
    ),
    23: (
        "Release history and audit records exist, but a documented periodic management review of all production changes is not evidenced. Recommended control is a monthly release-log review by the technical owner and a business approver, with quarterly Finance review of changes affecting financial reporting.",
        "Control gap",
        "Reviewed monthly release log and quarterly Finance sign-off.",
    ),
    30: (
        "Access is designed around unique accounts, Google Workspace SSO for workforce users, portal roles and workspace memberships, server-side authorization, Supabase row-level security, and separate administrator access at SaaS providers. A consolidated security policy covering all networks, devices, applications, databases, administrators, remote access, and service accounts is not evidenced.",
        "Partial control - policy required",
        "Approved information-security and access-control policies.",
    ),
    31: (
        "Workforce identity is primarily brokered through Google Workspace; external portal users use Supabase email/password. Password configuration is therefore split among providers. The portal does not currently require MFA, and the actual length, breached-password, lockout, recovery, and session settings must be exported and approved. Recommended policy is long unique passwords or SSO, breached-password blocking, no shared credentials, secure recovery, and MFA for privileged and financially sensitive access.",
        "Control gap - configuration evidence required",
        "Google, Supabase, QuickBooks, GitHub, and other provider authentication settings.",
    ),
    32: (
        "The intended policy is that employees use individual accounts and do not share credentials. However, non-personal demo accounts exist for executive testing and use a common test credential. Those accounts should be moved to an isolated demo tenant or blocked from production data and financial integrations. Management must issue and acknowledge a written prohibition on shared production credentials.",
        "Control gap",
        "Credential policy, employee acknowledgement, and demo-account isolation evidence.",
    ),
    33: (
        "Periodic forced password changes are not evidenced and should not be the primary control. Passwords should be changed when compromise is suspected, after recovery, when a shared or exposed credential is discovered, and when a service credential is rotated. Provider settings and the organization's final policy must be confirmed with SAX.",
        "Confirmation and policy decision required",
        "Approved password policy and provider configuration exports.",
    ),
    34: (
        "The portal adds active-user checks, least-privilege Viewer provisioning, role and organization scoping, server-side API checks, database row-level security, audit records, public-form challenge and rate limiting, and protected secret storage. Missing or unconfirmed layers include MFA for portal administrators, centralized identity alerts, periodic access review, and consolidated monitoring of privileged activity.",
        "Partial control - key gaps remain",
        "Security configuration, access-review evidence, and alerting/log review records.",
    ),
    35: (
        "Workforce portal access uses Google Workspace SSO through Supabase Auth. External Brand and Store users use the portal's Supabase password authentication. QuickBooks, Canix, GitHub, Supabase administration, hosting, banking, and payroll use their own provider standards unless separately federated. Management must confirm the current configuration for each in-scope application.",
        "Partially supported - system confirmation required",
        "Application-by-application authentication inventory.",
    ),
    38: (
        "Yes for the portal design: users authenticate through unique Google or Supabase identities, and protected server calls re-resolve the active user and role. Shared demo accounts are an exception and must remain isolated from live data. The same assertion for QuickBooks and other systems requires user-list evidence.",
        "Partially supported",
        "Portal user export and provider user/role listings.",
    ),
    39: (
        "The QuickBooks user and access-right listing has not been provided in the source material. The QuickBooks Primary Admin should export it directly from the production company, and management should review and sign it before submission.",
        "Open evidence request",
        "QuickBooks production user/role export with reviewer sign-off.",
    ),
    43: (
        "For UX OS, the current technical administrator manages the application and integrations, while portal Administrator roles control user and operational administration. QuickBooks, Google Workspace, Canix, GitHub, Supabase, hosting, payroll, and banking may have different administrators. Management must list the named administrator, title, backup, and approver for each in-scope system.",
        "Confirmation required",
        "Administrator register and screenshots/exports from each system.",
    ),
    44: (
        "The portal restricts privileged actions through Administrator roles, workspace permissions, server-side checks, and audit records. Whether any administrator also performs accounting functions in QuickBooks cannot be determined from the supplied evidence. Finance must identify such overlap and apply independent review where segregation is not practical.",
        "Partial control - segregation confirmation required",
        "QuickBooks role listing and segregation-of-duties assessment.",
    ),
    45: (
        "Portal role elevation and workspace assignments are controlled by an Administrator. First-time workforce SSO users default to Viewer. Management should formally require the user's manager or process owner to approve access and Finance/Controller approval for financially sensitive permissions. Approval practice for other systems must be confirmed.",
        "Partial control - approval workflow required",
        "Access request/approval samples and authority matrix.",
    ),
    46: (
        "Portal access changes are recorded, but a periodic management certification of roles across all in-scope systems is not evidenced. Recommended frequency is quarterly for administrators and financially sensitive access and at least semiannually for all other access.",
        "Control gap",
        "Completed access-review package and remediation evidence.",
    ),
    49: (
        "In UX OS, privileged functions require an active internal profile, the Administrator staff role or a specific permission, and a server-authorized request; protected data is not exposed merely because navigation is visible. Privileged access at QuickBooks and other SaaS providers is controlled in those platforms and must be separately evidenced.",
        "Supported for portal - provider evidence required",
        "Role/permission configuration and administrator exports.",
    ),
    50: (
        "Portal administrative changes are written to protected audit/event records. A recurring review of administrator activity is not yet evidenced. Recommended control is monthly review of privileged portal events and vendor administrator logs by someone other than the primary administrator, with exceptions documented and resolved.",
        "Partial control - recurring review required",
        "Monthly privileged-activity review and exception log.",
    ),
    54: (
        "Internal workforce users can be provisioned through Google Workspace SSO as least-privilege Viewers and are elevated only by an Administrator. Brand and Store users are invited or assigned to their approved organization/workspace. For other applications, the user request and approval process is not documented. A single joiner/mover/leaver form should identify requested system, role, scope, approver, and effective date.",
        "Partial control - enterprise process required",
        "New-user requests, approvals, and user listings.",
    ),
    55: (
        "An inactive portal profile is denied even if authentication succeeds. The missing control is a documented HR-to-IT termination notification with a same-day deadline and a checklist covering Google, QuickBooks, Canix, GitHub, Supabase, hosting, banking, payroll, email, service accounts, and devices. Actual termination samples must be tested.",
        "Partial control - termination workflow required",
        "Termination checklist, timestamps, and samples showing timely removal.",
    ),
    59: (
        "Recommended control: the user's manager or process owner approves the business need, and the designated system owner approves the role and scope before access is granted. Finance/Controller should approve QuickBooks and financially sensitive access. Current portal administration supports controlled assignment, but the management approval evidence must be formalized.",
        "Partial control - approval evidence required",
        "Approved access requests and role assignments.",
    ),
    61: (
        "Recommended control: all role, store, brand, and administrator changes require a documented request and owner approval before implementation. Elevation to Administrator should require executive approval and preferably a second approver. Current audit records support traceability, but the approval standard is not evidenced.",
        "Partial control - approval standard required",
        "Access-change requests, approvals, and audit records.",
    ),
    63: (
        "A periodic cross-system user-list review is not evidenced. Management should review administrators and financial access quarterly and all active users at least semiannually, compare the lists to current employees and approved external users, remove stale accounts, and retain reviewer sign-off.",
        "Control gap",
        "Reviewed user listings and documented removals.",
    ),
    69: (
        "Most core systems are SaaS and receive vendor-managed infrastructure patches. Custom UX OS code is versioned and release-tested. The organization still needs a documented monthly process for operating-system, browser, endpoint, dependency, container, and service patch status, with risk-based deadlines for critical vulnerabilities.",
        "Partial control - patch program required",
        "Patch policy, endpoint inventory, dependency alerts, and monthly compliance report.",
    ),
    71: (
        "No approved endpoint antivirus or EDR inventory was provided. Cloud vendors protect their hosted infrastructure under their own controls, but company laptops and any managed servers require centrally managed anti-malware/EDR, automatic updates, tamper protection, alert routing, and periodic coverage review. The portal's document scanner protects uploaded files but is not endpoint antivirus.",
        "Control gap",
        "Endpoint inventory, EDR/antivirus console export, update status, and exception list.",
    ),
    75: (
        "Recommended control: the technical owner reviews vendor notices, dependency/security alerts, operating-system and browser update status, and endpoint compliance monthly. Critical patches are evaluated promptly and overdue exceptions require documented risk acceptance. This consolidated review is not currently evidenced.",
        "Control gap",
        "Monthly patch review and remediation tickets.",
    ),
    77: (
        "Management should assess adequacy through coverage reports, signature/agent currency, alert testing, blocked-event review, exception tracking, and an annual configuration review. No such management review evidence was supplied.",
        "Control gap",
        "Quarterly EDR coverage report and annual control review.",
    ),
    79: (
        "No recurring independent vulnerability scan or penetration test was evidenced. Automated release and security-contract tests exist for the portal, but they are not a substitute for external vulnerability assessment. Recommended control is continuous dependency scanning, periodic authenticated vulnerability scanning, and an annual independent penetration test after material releases.",
        "Control gap",
        "Scan reports, remediation tickets, and penetration-test report.",
    ),
    81: (
        "Application and database platforms provide logs, and the portal keeps protected audit and integration events. A centralized SIEM or documented network-level monitoring process is not evidenced. Management should define which Google, GitHub, Supabase, hosting, endpoint, firewall, and vendor alerts are collected, who receives them, retention, severity thresholds, and response steps.",
        "Control gap",
        "Monitoring architecture, alert samples, log-retention settings, and incident tickets.",
    ),
    86: (
        "Core data is hosted by SaaS providers. The portal database and authentication data reside in Supabase; QuickBooks and Canix retain their authoritative records. Current documentation does not prove the enabled Supabase backup tier, configuration backups, export coverage, encryption, retention, or ownership. Management should document vendor backups and create a recoverable copy of critical portal configuration, mappings, and code.",
        "Control gap - vendor confirmation required",
        "Backup settings, retention, encryption, code repository status, and vendor documentation.",
    ),
    88: (
        "The actual backup frequency and type have not been evidenced. Recommended minimum is provider-managed continuous or daily database backup with point-in-time recovery where available, daily code/version control, and periodic encrypted exports of critical configuration and mappings. Recovery objectives should determine exact frequency and retention.",
        "Confirmation and design required",
        "Backup schedule, retention policy, RPO/RTO approval, and provider screenshots.",
    ),
    90: (
        "No documented periodic restore test was provided. Recommended control is a quarterly sample restore or recovery validation for portal data/configuration and an annual end-to-end recovery exercise, with results and exceptions retained.",
        "Control gap",
        "Restore test plan and completed test evidence.",
    ),
    92: (
        "A formal company disaster recovery and business continuity plan was not provided. Individual integrations include last-good snapshots and retry behavior, but those application features do not constitute an enterprise recovery plan.",
        "Control gap",
        "Approved disaster recovery and business continuity plan.",
    ),
    94: (
        "Not yet evidenced. The plan should name incident leadership and contacts, list material systems and dependencies, define priority and manual workarounds, assign recovery tasks, and approve recovery time and recovery point objectives for QuickBooks, Canix, Google Workspace, Supabase/UX OS, banking, payroll, and critical devices.",
        "Control gap",
        "DR plan with roles, contact tree, task order, system tiers, RTOs, and RPOs.",
    ),
    98: (
        "Backup success is not currently documented as a management control. Recommended evidence includes automated failure alerts, a daily or weekly dashboard review based on criticality, monthly owner sign-off, and periodic restore results rather than relying only on a vendor's successful-job indicator.",
        "Control gap",
        "Backup monitoring report, alert routing, and review sign-off.",
    ),
    100: (
        "No completed disaster-recovery test was provided. Recommended control is an annual tabletop covering loss of identity, portal/database outage, QuickBooks outage, Canix outage, compromised administrator access, and endpoint loss, followed by a technical restore test and tracked remediation.",
        "Control gap",
        "Exercise plan, attendance, results, recovery measurements, and remediation log.",
    ),
    106: (
        "The known systems are predominantly cloud-hosted, and no company-operated data center or computer room has been identified. If that is correct, vendor physical security should be assessed through SOC reports and contracts. Management must separately confirm how offices, networking equipment, and company devices are physically protected.",
        "Confirmation required",
        "Facility description, device inventory, and vendor SOC reports.",
    ),
    110: (
        "For vendor data centers, access is controlled by the vendors and should be supported by independent assurance reports. For company offices or network closets, management should document the actual control, such as locked doors, badge/key access, visitor procedures, and device screen lock and encryption. The current control was not provided.",
        "Confirmation required",
        "Physical-access procedure and photos or access configuration where appropriate.",
    ),
    112: (
        "Management should designate the office/facility owner who approves physical access and the system owner who approves access to any secure technology area. Vendor personnel access is governed by the vendor's control environment. Actual approvers must be named.",
        "Confirmation required",
        "Access-approval matrix and vendor assurance reports.",
    ),
    114: (
        "No entry-log evidence was supplied. If the entity has no controlled data center or computer room, state that clearly and describe the office or network-closet log that does exist. Vendor data-center logs are maintained by providers and tested through their assurance reports.",
        "Confirmation required",
        "Local entry logs or documented not-applicable rationale, plus vendor SOC evidence.",
    ),
    116: (
        "No periodic review evidence was supplied. If a local controlled room exists, management should define a quarterly review of access and entry exceptions. If not, the control can be documented as not applicable locally, with annual vendor assurance review covering hosted infrastructure.",
        "Confirmation required",
        "Quarterly access review or approved not-applicable assessment.",
    ),
}


def itgc_section(idx: int) -> str:
    if idx <= 9:
        return "IT governance and risk assessment"
    if idx <= 23:
        return "Change management"
    if idx <= 39:
        return "Logical security and passwords"
    if idx <= 50:
        return "Access management and privileged access"
    if idx <= 63:
        return "Joiners movers leavers and access review"
    if idx <= 81:
        return "Cybersecurity antivirus and patch management"
    if idx <= 100:
        return "Data backups and disaster recovery"
    return "Physical security"


def build_itgc() -> None:
    doc = Document()
    configure_document(
        doc,
        "IT General Controls Draft Responses and Remediation Plan",
        "Prepared for management review in response to the SAX questionnaire",
    )
    add_intro(
        doc,
        "This document answers the IT general-control questions using the current UX OS design, repository evidence, integration documentation, and decisions already made for the portal. It identifies where the evidence supports an auditor-facing response and where company-wide policy, named ownership, or third-party evidence is still required.",
        "UX OS has a strong technical foundation in server-side authorization, scoped access, protected secrets, audit records, read-only QuickBooks behavior, and controlled integration failure handling. The audit response is not yet ready to submit because several company-level controls are informal or unproven: MFA for privileged access, periodic user review, formal change approval, endpoint security, vulnerability testing, backup/restore evidence, disaster recovery, vendor assurance, and physical-security confirmation.",
    )

    doc.add_heading("Executive assessment", level=1)
    add_summary_table(
        doc,
        ["Domain", "Assessment", "Highest-priority next step"],
        [
            ["Governance", "Partially designed", "Name control owners and approve annual and quarterly review cadences."],
            ["Change management", "Technical controls exist; approvals are informal", "Require one change record linking approval, code, tests, deployment, and validation."],
            ["Access", "Strong portal authorization; enterprise evidence incomplete", "Enforce MFA for privileged access and perform an immediate cross-system user review."],
            ["Cybersecurity", "Application safeguards exist; endpoint and testing evidence missing", "Deploy or evidence managed EDR, patch review, vulnerability scanning, and alert handling."],
            ["Backup and recovery", "Not sufficiently evidenced", "Confirm provider backup tiers, approve RTO/RPO, and complete a restore and DR exercise."],
            ["Third parties and physical", "Cloud dependence understood; assurance not assembled", "Collect vendor SOC reports and document local physical controls or not-applicable rationale."],
        ],
        [1.45, 2.75, 2.65],
    )

    doc.add_heading("Systems likely in scope", level=1)
    add_bullets(
        doc,
        [
            "QuickBooks Online: accounting system of record for invoices, accounts receivable, payments, and the general ledger.",
            "Canix: authoritative physical and compliance inventory source that may affect inventory quantities and operational cutoff.",
            "UX OS and Supabase: portal identity, authorization, workflow, normalized snapshots, and read-only financial display; no QuickBooks posting route.",
            "Google Workspace: workforce identity, email, and SSO source.",
            "GitHub and the production hosting/deployment platform: source code, change history, and release delivery for UX OS.",
            "Payroll, banking, endpoint management, and any other applications that post to or support the general ledger: names and owners still need confirmation.",
            "Monday.com: historical and transitional workflow evidence; management must determine whether it remains financially relevant during the audit period.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("Questionnaire responses", level=1)
    current_section = None
    section_counter: Dict[str, int] = {}
    itgc_source = Document(SOURCE_ITGC)
    itgc_questions = {
        idx: " ".join(itgc_source.paragraphs[idx].text.split()).lstrip("• ").strip()
        for idx in ITGC_ANSWERS
    }
    for idx in sorted(itgc_questions):
        question = itgc_questions[idx]
        section = itgc_section(idx)
        if section != current_section:
            doc.add_heading(section, level=2)
            current_section = section
            section_counter[section] = 0
        section_counter[section] += 1
        answer, status, evidence = ITGC_ANSWERS[idx]
        add_qa(doc, f"IT-{idx:03d}", question, answer, status, evidence)

    doc.add_heading("Ninety day remediation plan", level=1)
    add_summary_table(
        doc,
        ["Priority", "Action", "Accountable owner", "Target", "Evidence of completion"],
        [
            ["P0", "Approve in-scope system register, owner RACI, and IT risk register.", "Executive management and technical owner", "0-15 days", "Approved register, risks, owners, and review calendar."],
            ["P0", "Export and review all privileged and financial-system users; remove stale access and document exceptions.", "System owners and Finance/Controller", "0-15 days", "Signed user lists and removal tickets."],
            ["P0", "Require MFA for Google Workspace, QuickBooks, GitHub, Supabase/hosting administrators, and portal administrators; isolate non-personal demo accounts from live data.", "Technical owner", "0-30 days", "Configuration exports, test results, and isolated demo-account evidence."],
            ["P0", "Adopt joiner/mover/leaver and quarterly privileged-access review procedures.", "Management, HR, and system owners", "0-30 days", "Approved procedure and first completed review."],
            ["P0", "Adopt change-management minimum: request, owner approval, Finance approval when applicable, tests, rollback, exact release identity, and post-release validation.", "Technical owner and business owners", "0-30 days", "Policy and two completed change samples."],
            ["P0", "Confirm backup settings, retention, encryption, and recovery objectives; complete the first restore test.", "Technical owner and system owners", "0-30 days", "Backup screenshots, RTO/RPO approval, and restore report."],
            ["P1", "Deploy or evidence managed endpoint EDR, encryption, screen lock, patching, and inventory.", "Technical owner", "31-60 days", "EDR and device-compliance exports."],
            ["P1", "Create vendor-risk register and collect current assurance for QuickBooks, Supabase, Google, Canix, GitHub, hosting, payroll, banking, email, and scanning providers.", "Technical owner and legal/management", "31-60 days", "Vendor register, contracts, SOC reports, and annual review."],
            ["P1", "Create one support and incident queue with severity, owner, response target, closure, and breach escalation.", "Technical owner and Operations", "31-60 days", "Procedure and sample ticket export."],
            ["P1", "Establish monthly patch/dependency review and centralized security-alert routing with retention.", "Technical owner", "31-60 days", "Monthly review and alert samples."],
            ["P2", "Complete an independent vulnerability assessment and penetration test and track remediation.", "Executive management and technical owner", "61-90 days", "Final report, remediation tickets, and retest results."],
            ["P2", "Approve the disaster recovery plan and complete a tabletop plus technical recovery exercise.", "Executive management, Finance, Operations, and technical owner", "61-90 days", "Plan, attendance, results, and remediation log."],
            ["P2", "Begin quarterly IT control review and annual management certification.", "Executive management", "By day 90 and recurring", "Signed quarterly control package and annual certification."],
        ],
        [0.55, 2.65, 1.35, 0.75, 1.55],
    )

    doc.add_heading("Evidence package to assemble for SAX", level=1)
    add_bullets(
        doc,
        [
            "In-scope system and administrator register with named owners and backups.",
            "QuickBooks, Google Workspace, Canix, GitHub, Supabase, hosting, payroll, and banking user/role exports.",
            "First quarterly access review and joiner/mover/leaver samples.",
            "Two production change samples linking approval, commit, test result, deployment, and validation.",
            "Current provider authentication settings and proof of MFA for privileged accounts.",
            "Endpoint inventory, encryption, patch, and antivirus/EDR coverage.",
            "Backup configuration, retention, restore-test evidence, and DR exercise results.",
            "Vendor assurance reports and annual review sign-off.",
            "Support/incident log and security monitoring evidence.",
        ],
    )

    doc.add_heading("Existing technical evidence reviewed", level=1)
    add_bullets(
        doc,
        [
            "Deployment readiness and Google Workspace SSO setup documentation.",
            "QuickBooks financial visibility and production-assessment documentation.",
            "Portal-native cutover, three-workspace architecture, and order-workflow documentation.",
            "Supabase database migrations, protected functions, row-level security, audit tables, and security-contract tests.",
            "Git source history and release verification scripts.",
        ],
    )
    doc.save(OUT_ITGC)


CYCLE_INFO = {
    "Revenue and accounts receivable": {
        "status": "Partially supported - Finance confirmation required",
        "systems": "QuickBooks Online is the accounting and accounts-receivable authority. UX OS owns portal orders and displays a scoped, read-only QuickBooks projection. Canix supports item and fulfillment facts. Signed agreements and approved prices are controlled outside or within approved portal records. UX OS does not automatically create a QuickBooks invoice.",
        "sources": "Signed customer agreements and amendments; approved price lists and discount approvals; portal orders and immutable order events; shipping, delivery, and return evidence; QuickBooks invoices, credit memos, customer statements, receipts, AR aging, reconciliations, and close support.",
        "owner": "Sales/Operations for order and delivery facts; Finance/Controller for invoices, revenue recognition, AR, credit, reconciliations, disclosures, and journal entries.",
    },
    "Cash receipts": {
        "status": "Finance confirmation required",
        "systems": "QuickBooks Online and the approved banking platform are expected to be authoritative. UX OS displays payment history from QuickBooks and does not collect money or write payment records.",
        "sources": "Bank statements, remittance advice, check images or deposit records where applicable, QuickBooks payment and deposit reports, customer ledger detail, cash reconciliations, and exception resolution.",
        "owner": "Finance/Controller and designated cash-receipt or accounting staff.",
    },
    "Provision for doubtful accounts": {
        "status": "Control design and Finance confirmation required",
        "systems": "QuickBooks AR aging and a controlled Finance analysis should support the reserve. UX OS may display aging information but is not the reserve calculation or posting system.",
        "sources": "AR aging, customer correspondence, collection history, subsequent receipts, write-off approvals, reserve calculation, management review, and related journal entries.",
        "owner": "Finance/Controller, with Sales input on customer facts.",
    },
    "Inventory control": {
        "status": "Partially supported - Operations and Finance confirmation required",
        "systems": "Canix is the physical and compliance inventory authority. UX OS reads Canix and retains the last successful complete snapshot; it does not edit Canix quantities. QuickBooks remains the financial ledger. The Canix-to-QuickBooks value reconciliation and several lot/cost-object decisions are not finalized.",
        "sources": "Canix item/package reports, manifests, transfer and adjustment records, production records, physical count sheets, variance approvals, consignment/ownership register, UX OS sync evidence, QuickBooks inventory/COGS schedules, and reconciliations.",
        "owner": "Operations for custody, counts, transfers, and Canix transactions; Finance/Controller for valuation and general-ledger reconciliation.",
    },
    "Inventory cost": {
        "status": "Open control design - Finance decision required",
        "systems": "Canix contains operational quantities and selected cost data, but the approved inventory-cost method and deterministic QuickBooks cost-object mapping are not evidenced. UX OS does not currently reconcile inventory value to the general ledger.",
        "sources": "Approved costing policy, bills of material, labor and overhead schedules, purchase invoices, freight/tax support, standard-cost or actual-cost calculations, variance analysis, valuation roll-forward, and reconciliation to QuickBooks.",
        "owner": "Finance/Controller, with Operations providing production, labor, yield, and inventory facts.",
    },
    "Inventory impairment": {
        "status": "Control design and Finance confirmation required",
        "systems": "Canix and UX OS can provide status, test result, age, facility, and quantity information. The reserve method, approval thresholds, and QuickBooks posting process were not provided.",
        "sources": "Inventory aging, test failures, quality holds, damage/obsolescence reports, expected selling prices, cost data, reserve calculation, approval, and write-down journal entry.",
        "owner": "Operations and Quality identify conditions; Finance/Controller estimates, approves, and records impairment.",
    },
    "Purchases and accounts payable": {
        "status": "Finance and Operations confirmation required",
        "systems": "QuickBooks Online is expected to be the accounts-payable and expense authority. Portal Brand purchase orders are operational records and do not automatically create QuickBooks transactions. The complete vendor approval, PO, receiving, and invoice process was not provided.",
        "sources": "Approved-vendor list, supplier onboarding and tax forms, purchase requisition/PO, approval, receiving report, supplier invoice, three-way match, supplier statement, AP aging, accrual, dispute record, and reconciliation.",
        "owner": "Requesting department and Operations for request/receipt; Finance/Controller for vendor master, coding, AP, accruals, reconciliations, and journal entries.",
    },
    "Cash payments and disbursements": {
        "status": "Finance and banking confirmation required",
        "systems": "QuickBooks and the approved banking or payment platform are authoritative. UX OS does not initiate or collect payments and does not modify accounting records.",
        "sources": "Approved invoice packet, payment batch, bank authorization, check/EFT detail, supplier ledger, bank statement, bank reconciliation, void/stop-payment evidence, and vendor-master change approval.",
        "owner": "Finance/Controller and designated banking approvers.",
    },
    "Property plant and equipment": {
        "status": "Finance confirmation required - not covered by UX OS",
        "systems": "QuickBooks and a controlled fixed-asset register should be authoritative. No UX OS module controls capitalization, depreciation, disposals, or fixed-asset safeguarding.",
        "sources": "Capital expenditure approval, invoice, receiving and placed-in-service evidence, fixed-asset register, location/custodian record, financing agreement, depreciation schedule, disposal approval, proceeds evidence, and reconciliations.",
        "owner": "Finance/Controller, with department custodians and executive approval under the capital policy.",
    },
    "Payroll": {
        "status": "HR and Finance confirmation required - payroll system not identified",
        "systems": "The payroll provider must be identified. QuickBooks records payroll expense and liabilities. UX OS does not process payroll, maintain pay rates, or make payroll payments.",
        "sources": "Approved offer/change/termination forms, tax and deduction elections, time records, bonus and overtime approvals, payroll register, funding and bank reports, tax filings, general-ledger posting, and payroll reconciliation.",
        "owner": "HR or executive management for employment data; payroll administrator for processing; Finance/Controller for funding, reconciliation, and general-ledger reporting.",
    },
    "New employees": {
        "status": "HR and management confirmation required",
        "systems": "The HR/payroll onboarding process and system were not provided. Portal access is a separate process and does not authorize payroll setup.",
        "sources": "Approved requisition or offer, background/eligibility records as applicable, tax forms, direct-deposit authorization, system-access request, and payroll setup review.",
        "owner": "Hiring manager and HR or executive management, with payroll administrator and system owners.",
    },
    "Employee terminations": {
        "status": "HR management and Finance confirmation required",
        "systems": "The HR/payroll termination process and system were not provided. Portal profiles can be deactivated, but an enterprise leaver checklist across payroll, Google, QuickBooks, Canix, GitHub, Supabase, hosting, banking, and devices is not yet evidenced.",
        "sources": "Termination notice, final-pay calculation, benefits/deduction changes, payroll removal evidence, system-access removal checklist, device return, and manager sign-off.",
        "owner": "Manager and HR or executive management, with payroll administrator and each system owner.",
    },
}


def financial_cycle(idx: int) -> str:
    if idx <= 54:
        return "Revenue and accounts receivable"
    if idx <= 81:
        return "Cash receipts"
    if idx <= 108:
        return "Provision for doubtful accounts"
    if idx <= 156:
        return "Inventory control"
    if idx <= 184:
        return "Inventory cost"
    if idx <= 203:
        return "Inventory impairment"
    if idx <= 249:
        return "Purchases and accounts payable"
    if idx <= 303:
        return "Cash payments and disbursements"
    if idx <= 396:
        return "Property plant and equipment"
    if idx <= 452:
        return "Payroll"
    if idx <= 461:
        return "New employees"
    return "Employee terminations"


def prefix(cycle: str, text: str) -> str:
    info = CYCLE_INFO[cycle]
    return f"{text} {info['owner']} Management must confirm the named personnel, frequency, thresholds, and retained evidence before submission."


def answer_financial_question(cycle: str, question: str) -> str:
    q = question.lower()
    info = CYCLE_INFO[cycle]
    if q.startswith("document the process for reconciling"):
        if cycle == "Inventory control":
            return prefix(cycle, "Recommended monthly process: reconcile Canix quantity reports and approved ownership classifications to the inventory valuation schedule, then reconcile the valuation schedule to the QuickBooks inventory and cost-of-goods-sold accounts. Investigate timing, unit, ownership, and cost differences before close. No automated reconciliation currently exists.")
        if cycle in {"Revenue and accounts receivable", "Purchases and accounts payable", "Cash receipts", "Cash payments and disbursements"}:
            return prefix(cycle, "Recommended monthly process: reconcile the detailed QuickBooks subledger or bank reconciliation to the corresponding general-ledger control account, investigate differences, retain support, and document independent review before close.")
        if cycle == "Payroll":
            return prefix(cycle, "Recommended process: reconcile the payroll register, payroll-provider funding, bank activity, payroll liabilities, and payroll expense posting to QuickBooks each pay period and again at month-end.")
        return prefix(cycle, "Recommended monthly process: reconcile the detailed supporting register or schedule to the related QuickBooks general-ledger account, investigate differences, and retain preparer and reviewer sign-off.")
    if q.startswith("list the names of the it applications"):
        return info["systems"]
    if q.startswith("list the types and titles of essential source documents"):
        return info["sources"]

    if cycle == "Revenue and accounts receivable":
        if "sales orders and contracts" in q and "created" in q:
            return prefix(cycle, "Sales orders are created in UX OS from an approved catalog and price snapshot. Contracts and amendments are retained as controlled signed records. New portal orders do not automatically create QuickBooks transactions; Finance creates the accounting record only after customer and order data are verified.")
        if "deliverables and performance obligations" in q:
            return prefix(cycle, "Deliverables and performance obligations should be identified from the executed contract, sales agreement, approved order, and shipping/service terms. Non-standard arrangements require Finance review before revenue recognition.")
        if "transaction prices" in q:
            return prefix(cycle, "Approved standard prices, contract or store prices, volume terms, promotions, and discounts are applied according to the approved commercial rules and frozen on the submitted order. Finance reviews variable consideration, returns, credits, taxes, and any non-standard term before invoicing and recognition.")
        if "granting credit" in q or "granted credit" in q or "credit limits" in q:
            return prefix(cycle, "Finance should approve credit status and limits using QuickBooks history, payment behavior, financial information, and management-approved exceptions. The portal must not infer an order hold solely from a QuickBooks balance; any hold or exception must be explicit and approved.")
        if "sales invoices" in q and ("generated" in q or "recorded" in q):
            return prefix(cycle, "Finance manually creates the QuickBooks invoice after verifying the customer, approved order, shipment or service evidence, price, discount, tax, and terms. UX OS has no invoice-creation route and displays only the later QuickBooks snapshot.")
        if "allocated to customer accounts" in q or "correct customer account" in q:
            return prefix(cycle, "Finance selects the verified QuickBooks customer identity and confirms the customer/store crosswalk. Display-name matching alone is not sufficient; exceptions remain unresolved until approved.")
        if "communicate with customers" in q or "account statements" in q:
            return prefix(cycle, "Finance communicates invoices, balances, statements, and payment status through approved QuickBooks/email processes. UX OS may display authorized invoice and payment history but does not collect payments.")
        if "resolving disputes" in q:
            return prefix(cycle, "Sales or Operations records the issue and supporting delivery/order facts; Finance evaluates any credit, return, or adjustment. Accounting changes require Finance approval and are entered only in QuickBooks, with the resolution retained.")
        if "subledgers reconcile" in q or "verify revenue" in q and "general ledger" in q:
            return prefix(cycle, "Finance reconciles the QuickBooks AR detail and revenue schedules to the general ledger, separately reconciles any deferred-revenue schedule, investigates differences, and documents reviewer approval before close.")
        if "deferred revenue" in q:
            return prefix(cycle, "Finance should identify prepayments or billed-but-unearned amounts from contract and fulfillment terms, maintain a deferred-revenue schedule, reconcile it to QuickBooks monthly, and approve recognition as obligations are satisfied. The actual process was not provided.")
        if "disclosures" in q:
            return prefix(cycle, "Finance prepares disclosures from the reconciled general ledger and supporting schedules and the Controller reviews classification, completeness, cutoff, significant judgments, and consistency with the financial statements.")
        if "delivery" in q or "deliveries" in q or "services performed" in q:
            return prefix(cycle, "Operations records shipment, delivery, receipt, exception, and supporting evidence against the approved order. Finance should invoice or recognize revenue only from authorized fulfillment evidence consistent with the contract.")
        if "prices approved" in q:
            return prefix(cycle, "Internal Sales, Operations, or Administrator roles approve published prices and discounts. Brand or store users may propose terms only within their scope; publication requires internal approval. Finance should approve any price rule that affects accounting or margin reporting.")
        if "separate deliverables" in q or "allocated to each element" in q:
            return prefix(cycle, "Finance reviews non-standard or bundled contracts, identifies each distinct obligation, documents the standalone selling price method, and approves the allocation before invoicing or recognition.")
        if "invoice" in q and ("accurate" in q or "accordance" in q or "products delivered" in q):
            return prefix(cycle, "Finance compares the invoice to the approved order, customer identity, price/discount terms, tax treatment, and shipment or service evidence. A completeness review compares fulfilled orders to invoices and open fulfillment at period end.")
        if "authorize invoices" in q:
            return prefix(cycle, "The invoice authorization matrix and limits were not provided. Recommended control is Finance preparation with Controller approval for non-standard, manual override, credit, or above-threshold invoices and independent review of invoices entered by privileged users.")
        if "recorded in the general ledger" in q or "correct amount" in q or "reconcile" in q:
            return prefix(cycle, "QuickBooks posts invoices and related entries to the ledger. Finance performs monthly AR aging, revenue, deferred-revenue, and general-ledger reconciliations; reviews cutoff and account coding; and resolves differences before close.")
        if "journal entries" in q:
            return "Journal entries are addressed in the separate journal-entry response. UX OS cannot create or post a QuickBooks journal entry. Finance must evidence unique access, support, independent approval, and monthly review."

    if cycle == "Cash receipts":
        if "methods" in q and "receive payments" in q:
            return prefix(cycle, "Payment methods are outside UX OS and must be confirmed by Finance. The response should list every actual method, such as ACH/wire, check, cash, or approved cannabis-banking channel, and identify which methods are permitted or prohibited.")
        if "customer deposits" in q:
            return prefix(cycle, "Customer deposits should be recorded as a liability or other approved account until the related performance obligation is met. Finance maintains and reconciles a deposit schedule. The actual account and process must be confirmed.")
        if "allocated to the correct customer" in q:
            return prefix(cycle, "Finance uses remittance details, invoice references, customer identity, and open items to apply receipts. Unapplied cash remains separately identified and is investigated rather than forced to a customer.")
        if "customer accounts" in q and "accurately" in q:
            return prefix(cycle, "Finance compares bank or deposit evidence, remittance detail, and the QuickBooks customer application, then investigates unapplied or misapplied receipts before close.")
        if "actual cash" in q or "cash accounts" in q or "reconcile" in q:
            return prefix(cycle, "Finance compares bank or deposit evidence to QuickBooks receipts and customer applications, investigates unmatched items, and completes an independent monthly bank and cash reconciliation.")
        if "recorded" in q and ("accounting system" in q or "general ledger" in q):
            return prefix(cycle, "An authorized Finance user records or imports the receipt in QuickBooks and applies it to the verified customer and invoice. Deposit and bank activity are then reconciled to QuickBooks; UX OS only reads the resulting payment history.")
        if "safeguard" in q or "limit" in q:
            return prefix(cycle, "Only named personnel should have bank, lockbox, check, deposit, or cash access. Banking credentials must be individual and protected by MFA; custody, recording, and reconciliation should be separated where staffing permits.")
        if "ensure receipts" in q:
            return prefix(cycle, "Finance compares bank or deposit evidence to QuickBooks receipts and customer applications, investigates unmatched items, and completes an independent monthly bank and cash reconciliation.")
        if "segregated duties" in q:
            return prefix(cycle, "Recommended segregation is: one person receives or initiates deposit information, another records or applies receipts, and an independent reviewer performs the bank reconciliation and reviews adjustments. Any staffing limitation requires documented compensating Controller review.")
        if "journal entries" in q:
            return "Journal entries are covered by the separate journal-entry response. Any cash adjustment must be supported and independently approved."
        return prefix(cycle, "Finance should document the end-to-end receipt, application, reconciliation, and review control in QuickBooks and the banking platform. UX OS is display-only and does not change cash or AR.")

    if cycle == "Provision for doubtful accounts":
        if "identified" in q and ("uncollectible" in q or "past due" in q):
            return prefix(cycle, "Finance reviews the QuickBooks AR aging, collection activity, disputes, payment history, subsequent receipts, and known customer conditions at least quarterly and at year-end. Sales provides current customer facts.")
        if "calculate the provision" in q or "appropriately estimated" in q:
            return prefix(cycle, "Finance should apply an approved methodology using aging categories, specific-risk accounts, historical loss experience, current conditions, and subsequent collections. The Controller reviews assumptions, calculations, and changes from the prior period.")
        if "write-offs" in q and ("recorded" in q or "accurately" in q):
            return prefix(cycle, "An approved write-off is entered in QuickBooks against the correct customer and allowance or expense account, with customer-level support and no deletion of the original invoice history.")
        if "policy" in q or "authorize write-offs" in q:
            return prefix(cycle, "The current write-off and reserve approval thresholds were not provided. Management should approve a matrix requiring Controller approval and executive approval above a defined threshold, with no self-approval by the preparer.")
        if "correct amount" in q or "general ledger" in q:
            return prefix(cycle, "Finance reconciles the reserve schedule, write-offs, AR aging, and related QuickBooks accounts at period end and reviews account, amount, and period before close.")
        if "journal entries" in q:
            return "Journal entries are covered by the separate journal-entry response; reserve and write-off entries require support and independent approval."

    if cycle == "Inventory control":
        if "physical counts" in q and "initiated" in q:
            return prefix(cycle, "Operations initiates an adjustment only after a documented count variance is recounted, investigated, and approved under a defined threshold matrix. The authorized user records the adjustment in Canix with reason and support.")
        if "transfers" in q and "initiated" in q:
            return prefix(cycle, "Authorized Operations users initiate facility, room, or custody transfers in Canix using the required compliance and manifest workflow. Source, destination, quantity, package/tag, and approver are retained.")
        if "movements" in q and "phases of production" in q:
            return prefix(cycle, "Operations records production consumption, conversion, split/merge, or output events in Canix using package and batch identities. Approved procedures and yields should support each movement.")
        if "purchases and purchase returns" in q:
            return prefix(cycle, "Inbound quantities and returns are recorded in Canix from approved receiving or return evidence. Finance separately records the accounting impact in QuickBooks and reconciles quantity and value.")
        if "sales and sales returns" in q:
            return prefix(cycle, "Canix records allocation/transfer/sale or return activity using package identity and compliance documents. Operations verifies fulfillment; Finance records the accounting transaction in QuickBooks.")
        if "adjusted in the accounting system" in q or "recorded in the accounting system" in q:
            return prefix(cycle, "Canix is the operational quantity system; QuickBooks is the financial ledger. No automatic value reconciliation is currently evidenced. Finance should use an approved monthly quantity-and-value bridge with documented exceptions.")
        if "consignment" in q or "third-party inventory" in q:
            return prefix(cycle, "Economic owner or arrangement type must be maintained separately from Brand. Third-party, toll, split, test, and company-owned inventory should be classified in an approved ownership register and excluded from company-owned inventory valuation where applicable. The ownership/lot control remains partly on hold.")
        if "track" in q and "safeguard" in q:
            return prefix(cycle, "Canix tracks items and packages by facility, location, package/tag, lot or batch fields, status, and source update. Physical access and movement authorization at each facility must be confirmed and documented by Operations.")
        if "monitor the existence" in q or "physical counts" in q and "accurate" in q:
            return prefix(cycle, "Operations should perform scheduled cycle counts and a year-end or risk-based full count using controlled count sheets, blind counts where practical, recounts of differences, tag/location completeness checks, and independent variance approval.")
        if "policy" in q and "authorize" in q:
            return prefix(cycle, "Management should approve role and dollar/quantity thresholds for adjustments, transfers, and production movements. Canix roles should restrict posting to authorized Operations users, and high-risk changes should receive independent review.")
        if "correct items" in q and ("shipped" in q or "received" in q):
            return prefix(cycle, "Operations compares item, package/tag, quantity, lot/batch, destination/source, license, and order/receiving document before confirming the movement. Exceptions are quarantined or held until resolved.")
        if "correct amount" in q or "general ledger" in q:
            return prefix(cycle, "Finance reconciles the approved inventory valuation schedule to QuickBooks and reviews ownership, cutoff, quantity, unit of measure, cost, account, and period. This bridge is not yet fully implemented and should be disclosed.")
        if "limit physical access" in q:
            return prefix(cycle, "Facility and room access should be limited to authorized personnel through locks or badges, visitor controls, inventory-room restrictions, and supervisory oversight. Actual site controls and access logs must be confirmed.")
        if "segregated duties" in q:
            return prefix(cycle, "Recommended segregation separates custody/counting, transaction entry, adjustment approval, valuation, and general-ledger reconciliation. Where staffing limits separation, the Controller and Operations manager should review all material adjustments and count variances.")
        if "journal entries" in q:
            return "Inventory-related journal entries are covered by the separate journal-entry response and require valuation support and independent approval."
        return prefix(cycle, "Operations controls the physical/compliance event in Canix; Finance controls the accounting consequence in QuickBooks. A documented monthly reconciliation is required because the systems do not share a single authoritative value record.")

    if cycle == "Inventory cost":
        if "journal entries" in q:
            return "Inventory-cost journal entries are covered by the separate journal-entry response and require approved cost schedules and independent review."
        if "basis or method" in q or "calculating inventory cost" in q:
            return prefix(cycle, "The approved costing basis was not provided. Finance must select and document the GAAP-consistent method for raw material, work in process, and finished goods, including unit-of-measure conversion, ownership, yield, and treatment of abnormal loss.")
        if "variances" in q:
            return prefix(cycle, "Finance should calculate purchase-price, usage/yield, labor, and overhead variances; investigate material differences; approve any capitalization or expense treatment; and reconcile the result to the ledger.")
        if "labor" in q or "overhead" in q:
            return prefix(cycle, "Finance should approve allocation pools, drivers, normal capacity, rates, exclusions, and update frequency. Operations supplies production volumes and labor/yield data; Finance tests the allocation and documents changes.")
        if "standard costs" in q:
            return prefix(cycle, "If standard cost is used, Finance maintains the standard-cost build, approves changes before use, compares standards to actual results, and reviews variances at least quarterly. Whether standard cost is currently used must be confirmed.")
        if "taxes" in q or "duties" in q or "shipping" in q or "other inventory costs" in q:
            return prefix(cycle, "Finance identifies capitalizable freight, taxes, duties, and conversion costs from invoices and policy, excludes selling/administrative and abnormal costs, and documents the allocation to inventory.")
        if "disclosures" in q:
            return prefix(cycle, "Finance prepares inventory disclosures from the approved costing policy and reconciled valuation schedule; the Controller reviews basis, classification, reserves, commitments, ownership, and consistency with the financial statements.")
        if "authorized" in q or "recorded" in q or "correct amount" in q:
            return prefix(cycle, "Every cost-model or rate change should have documented Finance approval, effective date, calculation, source data, version, and reconciliation to the inventory valuation and QuickBooks posting. This formal control is not yet evidenced.")
        return prefix(cycle, "The inventory-cost control cannot be represented as operating until Finance approves the costing policy, system mapping, and reconciliation evidence.")

    if cycle == "Inventory impairment":
        if "journal entries" in q:
            return "Impairment journal entries are covered by the separate journal-entry response and require the approved reserve analysis and independent review."
        if "slow moving" in q or "obsolete" in q or "damaged" in q or "below cost" in q:
            return prefix(cycle, "Operations and Quality should identify aged, failed, held, damaged, obsolete, recalled, or below-cost inventory from Canix/UX OS reports and physical inspection. Finance compares carrying cost to expected recoverable value and documents the conclusion.")
        if "calculate the provision" in q or "appropriately estimated" in q:
            return prefix(cycle, "Finance applies an approved reserve method using age, status, condition, test results, expected selling price, disposal/processing alternatives, and specific known losses. The Controller reviews assumptions and period-over-period changes.")
        if "write-down" in q:
            return prefix(cycle, "After approval, Finance records the write-down in QuickBooks against the correct inventory and expense/reserve accounts and preserves the Canix quantity and operational history unless an authorized operational disposal also occurs.")
        if "policy" in q or "authorize" in q:
            return prefix(cycle, "Management should approve reserve and write-down thresholds, reviewer independence, and escalation for material items. Actual limits were not provided.")
        if "correct amount" in q or "general ledger" in q:
            return prefix(cycle, "Finance reconciles the impairment schedule to inventory valuation and QuickBooks, verifies account and period, and retains the Controller's review before close.")

    if cycle == "Purchases and accounts payable":
        if "purchase orders generated" in q:
            return prefix(cycle, "The company purchasing process was not provided. Recommended practice is a numbered requisition/PO approved under an authority matrix before commitment. Portal Brand purchase orders are operational and do not create a QuickBooks payable.")
        if "authorizing suppliers" in q or "authorized suppliers" in q:
            return prefix(cycle, "Finance and the requesting department should approve suppliers after identity, tax, license/compliance, conflicts, bank-change, and business-need review. Payment terms are documented in the contract or vendor master and changes require independent approval.")
        if "outside of the purchase order" in q:
            return prefix(cycle, "Non-PO purchases should be limited to approved categories or emergencies, documented with business purpose and receipt, and approved under the same dollar authority before payment. Recurring exceptions should be reviewed.")
        if "goods received" in q or "services received" in q:
            return prefix(cycle, "Operations or the requester documents item/service, quantity, condition, date, and PO/contract reference. Finance records the purchase only from authorized evidence and investigates invoice-before-receipt or receipt-before-invoice items.")
        if "purchase invoices" in q or "accounts payable processed" in q:
            return prefix(cycle, "Finance validates the supplier, duplicate risk, invoice number/date, PO or approval, receipt, price, tax, account/class, and payment terms before recording the bill in QuickBooks.")
        if "invoice has not been received" in q or "goods or services have not been received" in q or "open purchase orders" in q:
            return prefix(cycle, "Finance and Operations review open POs, uninvoiced receipts, prepaid invoices, and unmatched bills at month-end and record or reverse accruals as appropriate.")
        if "statements from suppliers" in q:
            return prefix(cycle, "Finance compares supplier statements to the vendor ledger, investigates missing invoices, credits, duplicates, and unapplied payments, and documents resolution.")
        if "resolving disputes" in q:
            return prefix(cycle, "The requester or Operations confirms the factual issue; Finance places the bill or payment on hold and communicates with the supplier. Credits or adjustments require approval and are recorded in QuickBooks.")
        if "foreign currency" in q or "translated" in q:
            return prefix(cycle, "Finance should use the approved QuickBooks/company exchange-rate policy for initial recognition and period-end remeasurement. Management must confirm whether foreign-currency purchases exist; if none, document not applicable.")
        if "correct supplier account" in q:
            return prefix(cycle, "Finance selects the verified active vendor record and reviews duplicates, parent/alternate names, tax identity, and changes. Supplier bank or remittance changes require independent callback or equivalent verification.")
        if "in accordance with the terms" in q or "consistent with" in q:
            return prefix(cycle, "A three-way or appropriate two-way match compares the approved PO/contract, receipt or service confirmation, and invoice. Tolerance exceptions require documented approval.")
        if "recorded in the correct amount" in q or "subledgers agree" in q or "general ledger" in q:
            return prefix(cycle, "Finance reviews AP aging, accruals, unmatched items, vendor statements, cutoff, and account coding and reconciles the AP subledger to the QuickBooks general ledger monthly.")
        if "journal entries" in q:
            return "Purchase and AP journal entries are covered by the separate journal-entry response and require approved support and independent review."
        return prefix(cycle, "The complete AP workflow must be confirmed by Finance; UX OS does not replace QuickBooks vendor, bill, or payment controls.")

    if cycle == "Cash payments and disbursements":
        if "methods" in q and "make payments" in q:
            return prefix(cycle, "Finance must list actual payment methods, such as ACH/wire, check, cash, or approved payment platform, and identify who can initiate, release, void, and reconcile each method. UX OS does not make payments.")
        if "initiated" in q:
            return prefix(cycle, "A payment is initiated only from an approved invoice/payment packet and due-date review. The initiator creates the batch in the bank or payment system; an authorized approver independently releases it according to thresholds.")
        if "deposits and prepayments" in q:
            return prefix(cycle, "Finance records approved deposits/prepayments to the proper asset account, maintains a schedule, and reclassifies the amount when goods/services are received. The schedule is reconciled monthly.")
        if "allocated to the correct supplier" in q:
            return prefix(cycle, "Finance uses vendor identity, bill number, payment batch, and remittance detail to apply the payment. Unapplied or disputed items are separately identified and investigated.")
        if "supplier accounts" in q and "accurately" in q:
            return prefix(cycle, "Finance compares the approved payment batch and remittance detail to the QuickBooks vendor application and investigates unapplied, duplicate, or misapplied payments.")
        if "actual cash" in q or "cash accounts" in q or "agree to the general ledger" in q:
            return prefix(cycle, "An independent Finance reviewer completes monthly bank reconciliations, investigates outstanding and unusual items, and agrees the adjusted bank balance to QuickBooks before approval.")
        if "track payments" in q or "recorded in the accounting system" in q or "general ledger" in q:
            return prefix(cycle, "Finance records the payment in QuickBooks against the correct vendor and bill, retains bank/check confirmation, and reconciles cleared items to the bank statement and general ledger.")
        if "limit the ability" in q or "safeguarding" in q:
            return prefix(cycle, "Bank and payment access should be limited to named users with unique credentials and MFA. Initiation, approval, recording, and reconciliation should be separated, with transaction and daily limits.")
        if "authorize payments" in q or "in accordance with what was authorized" in q:
            return prefix(cycle, "Payment approval follows a written dollar authority matrix. The approver verifies the invoice packet, vendor, amount, due date, bank details, and duplicate risk before release. Privileged or above-threshold payments require an additional approver.")
        if "supplier" in q and "account information" in q:
            return prefix(cycle, "Vendor bank/remittance changes should require a documented request, independent verification using previously known contact information, segregation from payment release, and approval before the master record changes.")
        if "reconciles" in q:
            return prefix(cycle, "An independent Finance reviewer completes monthly bank reconciliations, investigates outstanding and unusual items, and agrees the adjusted bank balance to QuickBooks before approval.")
        if "segregated duties" in q:
            return prefix(cycle, "Recommended segregation separates vendor maintenance, purchase approval, payment initiation, payment release, accounting entry, and bank reconciliation. Staffing conflicts require independent Controller or executive review.")
        if "journal entries" in q:
            return "Payment-related journal entries are covered by the separate journal-entry response."
        return prefix(cycle, "The actual banking process and approvers were not provided and must be documented from Finance and bank evidence.")

    if cycle == "Property plant and equipment":
        if "disposals" in q or "disposing" in q:
            return prefix(cycle, "A custodian initiates a disposal request describing asset, reason, condition, proceeds, and method. Management approves before sale, scrap, transfer, or abandonment; Finance removes cost and accumulated depreciation from the register and ledger and records any approved gain or loss.")
        if "gains or losses" in q:
            return prefix(cycle, "Finance calculates proceeds less net book value, records the approved gain or loss, and agrees proceeds to bank activity and the disposal record.")
        if "depreciation expense" in q or "depreciation is calculated" in q or "depreciable" in q:
            return prefix(cycle, "Finance calculates depreciation from the controlled fixed-asset register using approved cost, useful life, method, residual value, and placed-in-service/disposal dates, reviews the calculation, and posts it to QuickBooks.")
        if "purchases of property" in q or "purchasing property" in q:
            return prefix(cycle, "Capital purchases should begin with a documented capital request describing purpose, vendor, amount, budget, useful life, and location, approved under the capital authority matrix before commitment.")
        if "recorded in the accounting system" in q or "posting property" in q:
            return prefix(cycle, "Finance records the asset in QuickBooks and the fixed-asset register from the approved invoice and placed-in-service evidence, assigning cost, class, location, custodian, useful life, method, and in-service date.")
        if "construction in progress" in q:
            return prefix(cycle, "Finance maintains a project-level CIP schedule, reconciles additions to QuickBooks, reviews completion status periodically, and transfers approved completed projects to fixed assets on the supported placed-in-service date.")
        if "interest" in q:
            return prefix(cycle, "If applicable, Finance identifies qualifying projects and borrowing costs, calculates capitalized interest under the approved policy, and obtains Controller review. If not applicable, management should document that conclusion.")
        if "financing arrangements" in q:
            return prefix(cycle, "Finance reviews the agreement, records the related asset and liability under the applicable accounting guidance, and reconciles principal, interest, and asset cost. Applicability must be confirmed.")
        if "track physical assets" in q or "inventory of property" in q:
            return prefix(cycle, "Assets should be tagged or uniquely identified by location and custodian. Departments confirm existence periodically, Finance investigates differences, and disposals or transfers require approval.")
        if "consistent with what was approved" in q or "not improperly expensed" in q:
            return prefix(cycle, "Finance compares purchases to the approved capital request and capitalization policy, reviews repairs/maintenance versus capital additions, and documents the account determination.")
        if "proceeds" in q or "misappropriated" in q:
            return prefix(cycle, "Sale proceeds should be paid directly to the company, deposited by authorized personnel, matched to the disposal, and independently reconciled to the bank and ledger.")
        if "useful lives" in q or "depreciation methods" in q:
            return prefix(cycle, "Management approves useful-life and depreciation-method policy by asset class. Exceptions or changes require Controller approval and documented accounting support.")
        if "correct amount" in q or "general ledger" in q or "transactions are recorded correctly" in q:
            return prefix(cycle, "Finance reconciles the fixed-asset, CIP, accumulated-depreciation, and depreciation schedules to QuickBooks monthly or quarterly and reviews additions, disposals, classification, cutoff, and period.")
        if "journal entries" in q:
            return "Fixed-asset journal entries are covered by the separate journal-entry response and require the approved fixed-asset schedule."
        return prefix(cycle, "The fixed-asset process is outside UX OS and must be completed from Finance's QuickBooks and fixed-asset evidence.")

    if cycle == "Payroll":
        if "pay rates" in q or "deductions" in q:
            return prefix(cycle, "A manager or HR initiates the approved change in writing; a separate payroll administrator enters it; and a reviewer compares the next payroll register to the approval. Tax elections and employee-authorized deductions follow provider controls.")
        if "overtime" in q or "bonuses" in q or "other payments" in q:
            return prefix(cycle, "The employee's manager initiates and approves overtime or bonus information under the compensation policy. Above-threshold or executive payments require additional approval before payroll entry.")
        if "hours worked" in q:
            return prefix(cycle, "Employees record time in the approved timekeeping method and supervisors approve it before payroll processing. The system and actual review cadence must be identified.")
        if "changes in employment status" in q:
            return prefix(cycle, "HR or authorized management provides approved hire, leave, and termination data to payroll. The payroll administrator enters the change and an independent reviewer compares the payroll register to the approved change list.")
        if "calculating and preparing payroll" in q or "payroll calculations" in q:
            return prefix(cycle, "The payroll provider calculates gross pay, deductions, taxes, and net pay from approved master data and time/compensation inputs. Payroll and Finance review variance and exception reports before release.")
        if "making payroll payments" in q or "payroll payments are authorized" in q:
            return prefix(cycle, "Payroll funding and release require authorized review of the final payroll register, employee count, gross-to-net totals, bank account, and funding amount. Access must be limited and protected by MFA.")
        if "adjustments" in q or "inaccuracies corrected" in q:
            return prefix(cycle, "Payroll corrections require documented cause, recalculation, employee impact, tax treatment, and approval. Off-cycle payments or reversals receive separate approval and are reconciled.")
        if "recorded in the general ledger" in q or "correct amount" in q or "cash accounts" in q:
            return prefix(cycle, "Finance reconciles the payroll register, provider funding, tax/benefit liabilities, bank activity, and QuickBooks posting each payroll and month-end, then resolves differences before close.")
        if "safeguard" in q or "segregated duties" in q:
            return prefix(cycle, "Payroll, banking, personnel data, and reports should be restricted to named users. HR authorization, payroll entry, payroll release, accounting, and reconciliation should be separated or subject to independent compensating review.")
        if "cash" in q and "unclaimed" in q:
            return prefix(cycle, "Management should confirm whether cash payroll exists. Recommended policy is electronic or controlled check payment; any unclaimed check is logged, secured, voided or reissued under approval, and never treated as unrecorded cash.")
        if "journal entries" in q:
            return "Payroll journal entries are covered by the separate journal-entry response and require the approved payroll register and reconciliation."
        return prefix(cycle, "The payroll provider, actual roles, and current operating procedure must be identified before this response can be finalized.")

    if cycle == "New employees":
        if "hiring new employees" in q:
            return prefix(cycle, "A manager initiates an approved requisition or offer; HR or authorized management completes eligibility and onboarding; payroll enters the employee only from approved documentation; and system access follows a separate least-privilege request.")
        if "authorized" in q or "entered into the payroll" in q:
            return prefix(cycle, "The payroll administrator should add only employees supported by an approved offer/start record, tax and payment forms, and manager/HR authorization. A reviewer compares new employees on the first payroll register to the approved hire list.")

    if cycle == "Employee terminations":
        if "notifying the payroll department" in q:
            return prefix(cycle, "The manager or HR provides payroll and system owners a dated termination notice before the last day when possible, including final-pay, benefit/deduction, device return, and access-removal requirements.")
        if "removed from the payroll system" in q:
            return prefix(cycle, "Payroll deactivates the employee after final-pay processing, and a reviewer compares terminated employees to the next payroll register. The enterprise leaver checklist also removes application, email, banking, and device access on the effective date.")

    return prefix(cycle, "The current operating procedure was not included in the supplied materials. The recommended control should be documented, assigned, performed, and retained as evidence.")


def build_financial_controls() -> None:
    doc = Document()
    configure_document(
        doc,
        "Financial Process Internal Controls Draft Responses",
        "Prepared for management review in response to the SAX questionnaire",
    )
    add_intro(
        doc,
        "This document provides a complete response framework for the financial-cycle questionnaire. It incorporates known UX OS, QuickBooks, and Canix boundaries and supplies recommended auditor-facing language where the underlying business process was not provided.",
        "The technology boundary is clear: QuickBooks remains the accounting authority, Canix remains the physical/compliance inventory authority, and UX OS is a controlled workflow and read-only projection layer for selected financial data. Most accounting-cycle responses still require Finance, Operations, HR, or management to confirm the people, frequency, thresholds, and evidence. Inventory costing and several cross-system reconciliations remain material open design items.",
    )

    doc.add_heading("Cycle readiness summary", level=1)
    rows = []
    for cycle in (
        "Revenue and accounts receivable",
        "Cash receipts",
        "Provision for doubtful accounts",
        "Inventory control",
        "Inventory cost",
        "Inventory impairment",
        "Purchases and accounts payable",
        "Cash payments and disbursements",
        "Property plant and equipment",
        "Payroll",
    ):
        info = CYCLE_INFO[cycle]
        rows.append([cycle, info["status"], info["owner"]])
    add_summary_table(doc, ["Cycle", "Readiness", "Proposed accountable roles"], rows, [2.0, 2.25, 2.75])

    doc.add_heading("Important system boundaries", level=1)
    add_bullets(
        doc,
        [
            "UX OS does not automatically create QuickBooks invoices, payments, bills, journal entries, or other general-ledger transactions.",
            "QuickBooks is authoritative for accounting records; the portal reads a limited normalized snapshot.",
            "Canix is authoritative for physical/compliance inventory quantities and operational status, while inventory valuation and ledger reconciliation remain Finance controls.",
            "A field match or portal display is not, by itself, evidence that the subledger agrees to the general ledger.",
            "Payroll, banking, and fixed-asset applications and owners must be named before the related sections are submitted.",
        ],
    )

    doc.add_page_break()
    current_cycle = None
    cycle_q_no = 0
    for idx, question in source_questions(SOURCE_IC):
        cycle = financial_cycle(idx)
        if cycle != current_cycle:
            if current_cycle is not None:
                doc.add_page_break()
            doc.add_heading(cycle, level=1)
            info = CYCLE_INFO[cycle]
            p = doc.add_paragraph()
            lead = p.add_run("System boundary. ")
            set_run_font(lead, bold=True)
            set_run_font(p.add_run(info["systems"]))
            set_paragraph_spacing(p, after=5, line=1.12)
            p = doc.add_paragraph()
            lead = p.add_run("Essential evidence. ")
            set_run_font(lead, bold=True)
            set_run_font(p.add_run(info["sources"]))
            set_paragraph_spacing(p, after=5, line=1.12)
            p = doc.add_paragraph()
            lead = p.add_run("Proposed accountable roles. ")
            set_run_font(lead, bold=True)
            set_run_font(p.add_run(info["owner"]))
            set_paragraph_spacing(p, after=8, line=1.12)
            current_cycle = cycle
            cycle_q_no = 0
        cycle_q_no += 1
        answer = answer_financial_question(cycle, question)
        evidence = CYCLE_INFO[cycle]["sources"]
        add_qa(doc, f"{cycle_q_no:02d}", question, answer, CYCLE_INFO[cycle]["status"], evidence)

    doc.add_page_break()
    doc.add_heading("Cross-cycle management confirmations", level=1)
    add_bullets(
        doc,
        [
            "Name the Finance/Controller, preparers, reviewers, banking approvers, payroll administrator, Operations inventory owners, and HR owner.",
            "Provide the approved authorization matrix for sales, pricing, credit, invoices, write-offs, purchases, payments, journal entries, inventory adjustments, capital expenditures, payroll, and administrator access.",
            "Confirm the monthly close calendar and supply signed reconciliations for cash, AR, AP, inventory, payroll, fixed assets, and significant estimates.",
            "Identify the payroll provider, banking platforms, fixed-asset register, and any other systems that affect the general ledger.",
            "Approve the inventory costing method, ownership treatment, impairment method, cost-object mapping, and Canix-to-QuickBooks reconciliation.",
            "Provide samples for each major process showing authorization, processing, review, and final accounting evidence.",
        ],
    )
    doc.add_heading("Highest priority control design items", level=1)
    add_summary_table(
        doc,
        ["Priority", "Item", "Why it matters", "Proposed owner"],
        [
            ["P0", "Inventory costing and reconciliation", "Current sources do not establish an approved valuation method or a complete Canix-to-QuickBooks value bridge.", "Finance/Controller and Operations"],
            ["P0", "Journal-entry policy and monthly close evidence", "SAX requires support for authorization, accuracy, period, account, and independent review.", "Finance/Controller"],
            ["P0", "QuickBooks user and role review", "Privileged access and segregation must be evidenced directly from the production company.", "QuickBooks Primary Admin and Finance"],
            ["P1", "Revenue and AP process narratives", "The portal boundary is known, but the actual invoice, bill, credit, receipt, and payment procedures require owner confirmation.", "Finance, Sales, and Operations"],
            ["P1", "Allowance and impairment methodologies", "Estimates need documented inputs, assumptions, review, and period-end posting controls.", "Finance/Controller"],
            ["P1", "Payroll and fixed-asset controls", "The systems, owners, and operating evidence were not provided.", "HR, Finance, and asset custodians"],
        ],
        [0.55, 1.8, 2.95, 1.7],
    )
    doc.save(OUT_IC)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for source in (SOURCE_JE, SOURCE_IC, SOURCE_ITGC):
        if not source.exists():
            raise FileNotFoundError(source)
    build_journal_entries()
    build_financial_controls()
    build_itgc()
    print(OUT_JE)
    print(OUT_IC)
    print(OUT_ITGC)


if __name__ == "__main__":
    main()

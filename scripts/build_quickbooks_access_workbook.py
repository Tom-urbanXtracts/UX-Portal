from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("output/quickbooks-access-controls-2026-10-07")
OUTFILE = OUTPUT / "QuickBooks Access Safeguards Implementation Workbook.docx"

NAVY = "17324D"
BLUE = "DDEBF7"
PALE_BLUE = "EEF5FA"
PALE_GOLD = "FFF2CC"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "6B7280"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"
GREEN = "E2F0D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_col_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def font_run(run, size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=5, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(doc, text, bold_lead=None, italic=False, after=6):
    p = doc.add_paragraph()
    style_paragraph(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        font_run(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        font_run(r, italic=italic)
    else:
        r = p.add_run(text)
        font_run(r, italic=italic)
    return p


def add_bullet(doc, text, level=0, checkbox=False):
    p = doc.add_paragraph(style="List Bullet" if not checkbox else None)
    p.paragraph_format.left_indent = Inches(0.28 + 0.22 * level)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    prefix = "[ ] " if checkbox else ""
    r = p.add_run(prefix + text)
    font_run(r)
    return p


def add_numbered_steps(doc, steps):
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{i}. ")
        font_run(r, bold=True)
        r = p.add_run(step)
        font_run(r)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        font_run(r, size=16 if level == 1 else 12, bold=True, color=BLACK)
    return p


def add_table(doc, headers, rows, widths, header_fill=NAVY, font_size=9.1, response_cols=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    keep_row_together(hdr)
    for i, (cell, header, width) in enumerate(zip(hdr.cells, headers, widths)):
        set_col_width(cell, width)
        set_cell_margins(cell)
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0)
        r = p.add_run(header)
        font_run(r, size=9, bold=True, color=WHITE)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        for col_idx, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            set_col_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if response_cols and col_idx in response_cols:
                set_cell_shading(cell, PALE_GOLD)
            elif row_idx % 2:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.03)
            r = p.add_run(str(value))
            font_run(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_response_block(doc, label, lines=2, prompt=None):
    if prompt:
        add_text(doc, prompt, after=3)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    keep_row_together(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GOLD)
    set_cell_margins(cell, top=100, start=120, bottom=100, end=120)
    p = cell.paragraphs[0]
    style_paragraph(p, after=2)
    r = p.add_run(label)
    font_run(r, size=9.5, bold=True, color=MID_GRAY)
    for _ in range(lines):
        p = cell.add_paragraph("________________________________________________________________________________")
        style_paragraph(p, after=2)
        for r in p.runs:
            font_run(r, size=9, color="A6A6A6")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_role_section(doc, number, name, purpose, allow, deny, steps):
    add_heading(doc, f"{number} {name}", level=2)
    add_text(doc, f"Purpose. {purpose}", bold_lead="Purpose. ")
    p = doc.add_paragraph()
    style_paragraph(p, after=3)
    r = p.add_run("Allow")
    font_run(r, bold=True)
    for item in allow:
        add_bullet(doc, item)
    p = doc.add_paragraph()
    style_paragraph(p, after=3)
    r = p.add_run("Do not allow")
    font_run(r, bold=True)
    for item in deny:
        add_bullet(doc, item)
    add_text(doc, "How to configure", bold_lead="How to configure", after=3)
    add_numbered_steps(doc, steps)
    add_table(
        doc,
        ["Completion item", "Your response"],
        [
            ["Role created", "[ ] Yes   [ ] No   [ ] Not applicable"],
            ["Exact QuickBooks role name", ""],
            ["Assigned users", ""],
            ["Configured by and date", ""],
            ["Notes or limitation", ""],
        ],
        [2.15, 4.85],
        font_size=9.4,
        response_cols={1},
    )


def build():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    for style_name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    # Keep the cover title clean. Some Word templates add a bottom border to
    # the built-in Title style, which reads like an unintended divider.
    title_style_ppr = styles["Title"].element.get_or_add_pPr()
    title_style_border = title_style_ppr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_ppr.remove(title_style_border)

    title = doc.add_paragraph(style="Title")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(7)
    r = title.add_run("QuickBooks Access Safeguards Implementation Workbook")
    font_run(r, size=25, bold=True)
    sub = doc.add_paragraph()
    style_paragraph(sub, after=14)
    r = sub.add_run("urbanXtracts  |  Management working document  |  7 October 2026")
    font_run(r, size=10.5, bold=True, color=MID_GRAY)

    add_text(
        doc,
        "This workbook is for reviewing and correcting QuickBooks user access before the SAX audit response is finalized. Complete the yellow fields, retain the requested evidence, and return the same file for final formatting and incorporation into the audit package.",
        after=7,
    )
    add_text(
        doc,
        "Primary finding. QuickBooks currently shows one Primary Admin, seven Company Admins, and one In-house Accountant. The recommended target is one Primary Admin, one named backup administrator, and restricted roles for everyone else.",
        bold_lead="Primary finding. ",
        after=7,
    )
    add_text(
        doc,
        "Implementation rule. Create and test restricted roles before reducing anyone's access. Do not share credentials, delete historical users, or remove the Primary Admin without an approved replacement and retained evidence.",
        bold_lead="Implementation rule. ",
        after=10,
    )

    add_heading(doc, "How to use this workbook", 1)
    for item in [
        "Review the current-access table and correct any inaccurate name, email, status, or role.",
        "Answer the responsibility questions before assigning roles.",
        "Create the recommended roles in QuickBooks using the step-by-step instructions.",
        "Test each role before reassigning a live user.",
        "Reassign users, enable two-step verification, and configure the monthly and quarterly reviews.",
        "Save the evidence listed in the final section and complete the management sign-off.",
        "Return the completed file for final formatting and reconciliation to the SAX questionnaires.",
    ]:
        add_bullet(doc, item, checkbox=True)

    add_heading(doc, "Current QuickBooks access snapshot", 1)
    add_text(doc, "Confirm this table against QuickBooks before making changes. The snapshot is based on the screenshots provided on 7 October 2026.")
    users = [
        ["Reiko Farmer", "reiko@urbanxtracts.com", "Primary admin", "Active", "Confirm"],
        ["Steven Felice", "steven.f@urbanxtracts.com", "Company admin", "Active", "Reassign"],
        ["Amrit Kharas", "amrit@urbanxtracts.com", "Company admin", "Active", "Reassign"],
        ["Janvi Poptani", "janvi@urbanxtracts.com", "Company admin", "Active", "Reassign"],
        ["Albert Schulman", "Albert@urbanxtracts.com", "Company admin", "Active", "Reassign"],
        ["Omeed Turan", "omeed@urbanxtracts.com", "Company admin", "Active", "Reassign"],
        ["Eran Sherin", "eran@urbanxtracts.com", "Company admin", "Active", "Reassign"],
        ["Tom Jalallar", "tom@urbanxtracts.com", "Company admin", "Active", "Reassign or backup admin"],
        ["Fred Sheinblum", "fsheinblum@gmail.com", "In house accountant", "Active", "Confirm engagement and access"],
    ]
    add_table(doc, ["User", "Email", "Current role", "Status", "Required review"], users, [1.25, 2.05, 1.3, 0.65, 1.75], font_size=8.4)
    add_response_block(doc, "Corrections to the current access snapshot", lines=3)

    add_heading(doc, "Responsibility decisions before changing access", 1)
    decisions = [
        ("1", "Who will remain the Primary Admin?", "The legally authorized account owner or executive responsible for the QuickBooks company."),
        ("2", "Who will be the single backup Company Admin?", "Choose one named person who can restore access and administer users when the Primary Admin is unavailable."),
        ("3", "Who creates bills?", "Person entering supplier invoices and supporting documents."),
        ("4", "Who approves bills?", "Person checking business purpose, receipt, price, coding, and authorization."),
        ("5", "Who releases payments?", "Person sending an already approved payment through QuickBooks or the banking platform."),
        ("6", "Who creates or changes vendors?", "Identify the person maintaining vendor identity, address, tax, and remittance information."),
        ("7", "Who independently verifies vendor banking changes?", "This should not be the same person requesting the change or releasing the payment."),
        ("8", "Who creates invoices and manages accounts receivable?", "Identify responsibility for customers, invoices, credits, and receipt application."),
        ("9", "Who reconciles bank and credit-card accounts?", "The reconciler should not release the same payments whenever staffing permits."),
        ("10", "Who prepares journal entries?", "Identify standard, recurring, non-standard, period-end, correction, and management-entry preparers."),
        ("11", "Who reviews and approves journal entries?", "Reviewer should be independent of the preparer and retain evidence."),
        ("12", "Who runs and reviews payroll?", "Name the payroll preparer, approver, funding approver, and reconciler."),
        ("13", "Who needs report-only access?", "Executives and managers who need visibility but do not enter transactions."),
    ]
    for number, question, instruction in decisions:
        add_heading(doc, f"Decision {number}  {question}", 2)
        add_text(doc, instruction, after=3)
        add_response_block(doc, "Your answer", lines=2)

    add_heading(doc, "Recommended QuickBooks roles", 1)
    add_text(doc, "Use these as the starting role set. If QuickBooks does not offer a required granular permission, document the limitation and use the compensating review described in this workbook.")

    common_steps = [
        "Open Settings and select Manage users.",
        "Select the Roles tab and then Add role.",
        "Enter the role name and description shown below.",
        "Enable only the listed access. Leave unrelated areas disabled.",
        "Review the permission summary and save the role.",
        "Assign the role to a test user and confirm the permitted and prohibited actions.",
        "Record the test result and only then assign the role to live users.",
    ]

    role_specs = [
        ("1", "Primary Admin", "Maintain legal ownership and complete account recovery authority.",
         ["Full account administration", "Primary-admin ownership and recovery", "User and subscription administration"],
         ["Routine bill entry", "Routine payment release", "Routine journal preparation", "Routine bank reconciliation"],
         ["Confirm the current Primary Admin is still authorized by management.", "If a transfer is needed, first make the incoming person a Company Admin.", "From Manage users, use the action menu to make the approved person Primary Admin.", "Retain the approval and before-and-after screenshots."]),
        ("2", "Backup Company Admin", "Provide emergency administrative continuity without creating a large administrator population.",
         ["User administration", "Role administration", "Emergency account recovery and configuration"],
         ["Routine accounting transactions", "Routine vendor maintenance", "Routine payment release", "Routine bank reconciliation"],
         common_steps),
        ("3", "Finance Controller", "Control accounting, close, journals, reconciliations, and financial reporting without user administration.",
         ["Accounting and chart-of-accounts functions", "Journal entries and recurring entries", "Reconciliations", "Financial reports", "Close and review functions required by policy"],
         ["Manage users", "Change administrator roles", "Routine bill preparation and payment release when segregation is available", "Payroll unless separately assigned"],
         common_steps),
        ("4", "Bookkeeper", "Perform routine accounting entry without administrator or payment-release authority.",
         ["Routine sales and expense transactions", "Customer and vendor records as approved", "Bank-feed classification", "Standard operational reports"],
         ["Manage users", "Change company settings", "Release payments", "Approve own transactions", "Close books or override closed periods"],
         common_steps),
        ("5", "Accounts Receivable", "Manage customer transactions without access to vendors, bills, payroll, or administrative settings.",
         ["Customers", "Invoices and sales receipts", "Approved credit memos", "Receipt application", "Accounts-receivable reports"],
         ["Vendors and bills", "Payment release", "Payroll", "Journal entries", "User administration", "Bank reconciliation"],
         common_steps),
        ("6", "Accounts Payable Entry", "Enter supplier bills while keeping approval and payment release separate.",
         ["View approved vendors", "Create bills", "Attach supporting documents", "View bill status"],
         ["Approve bills", "Release payments", "Reconcile bank accounts", "Delete approved bills", "Change vendor banking information unless separately controlled"],
         common_steps),
        ("7", "Bill Approver", "Approve properly supported bills without creating or paying them.",
         ["View bill and support", "Approve or reject bills", "Record an approval note"],
         ["Create bills", "Edit vendors", "Release payments", "Reconcile bank accounts", "User administration"],
         ["Use the predefined Bill approver role if its permission summary matches this definition.", "Open Settings, Manage users, Users, and edit the approved user.", "Select Bill approver and review the displayed permissions.", "Save and test that the user cannot create or pay a bill.", "Retain the test evidence."]),
        ("8", "Payment Releaser", "Release only previously approved payments while preventing bill creation and vendor maintenance.",
         ["View approved bills", "Release approved payments", "View payment status and remittance"],
         ["Create or approve bills", "Create or change vendors", "Change vendor banking information", "Reconcile bank accounts", "User administration"],
         common_steps),
        ("9", "Inventory Specialist", "Maintain products, services, and inventory records without financial-administration access.",
         ["Products and services", "Inventory quantities and related operational records", "Inventory reports needed for the role"],
         ["Banking", "Payroll", "Journal entries", "Bill payment", "User administration"],
         common_steps),
        ("10", "Payroll Manager", "Run approved payroll while limiting general accounting and administrative access.",
         ["Payroll preparation and payroll reports", "Approved employee payroll fields", "Payroll tax and filing functions required by the role"],
         ["User administration", "Unrelated banking", "Vendor payments", "General journals outside approved payroll posting", "Company settings"],
         common_steps),
        ("11", "Executive Reports", "Provide management visibility without transaction or administrative access.",
         ["Approved company financial reports", "Management dashboards and budgets where appropriate"],
         ["Create, edit, delete, approve, or pay transactions", "Payroll details unless specifically authorized", "User administration", "Vendor banking information"],
         ["Use View company reports when its permission summary is appropriate.", "Open Settings, Manage users, Users, and edit the approved user.", "Select View company reports.", "Confirm that payroll and contact information remain restricted.", "Save and test the role."]),
        ("12", "External Accountant", "Give the accountant the accounting access needed for the engagement without routine company administration.",
         ["Accounting records and reports required for the engagement", "Reconciliations and journal review as contractually approved"],
         ["User administration unless explicitly authorized", "Payroll unless in scope", "Payment release", "Unrestricted vendor maintenance"],
         ["Confirm the accountant's active engagement and required scope.", "Use the accountant invitation or In-house Accountant role appropriate to the subscription.", "Review the resulting permission summary before sending the invitation.", "Require individual credentials and two-step verification.", "Review the access at least quarterly and remove it promptly when the engagement ends."]),
    ]

    for spec in role_specs:
        add_role_section(doc, *spec)

    add_heading(doc, "User reassignment worksheet", 1)
    add_text(doc, "Complete one row for every active user. A user should receive the minimum role necessary for current responsibilities. Do not use Company Admin as a convenience role.")
    assignment_rows = []
    for name, email, current, _, _ in users:
        suggested = "Keep Primary Admin if authorized" if current == "Primary admin" else ("Confirm In-house Accountant" if current == "In house accountant" else "Select restricted role")
        assignment_rows.append([name, current, suggested, "", "", "[ ]"])
    add_table(
        doc,
        ["User", "Current role", "Starting recommendation", "Final role", "Approved by and date", "Done"],
        assignment_rows,
        [1.25, 1.15, 1.65, 1.25, 1.35, 0.35],
        font_size=7.8,
        response_cols={3, 4, 5},
    )
    add_response_block(doc, "Users to deactivate or remove from active access", lines=3)
    add_response_block(doc, "Temporary access and planned expiration dates", lines=3)

    add_heading(doc, "How to change an existing user's role", 1)
    add_numbered_steps(doc, [
        "Open Settings and select Manage users.",
        "Select the Users tab.",
        "Locate the user and select Edit in the Action column.",
        "Choose the approved predefined or custom role.",
        "Review the permission summary shown by QuickBooks.",
        "Save the change.",
        "Ask the user to sign out and sign back in, then test the expected actions.",
        "Capture the updated Users screen and file it with the approval evidence.",
    ])

    add_heading(doc, "How to deactivate a user", 1)
    add_numbered_steps(doc, [
        "Confirm the employee, contractor, or advisor no longer requires access and retain the termination or removal approval.",
        "Open Settings and select Manage users.",
        "Locate the user and open the Action menu.",
        "Select Delete or deactivate, depending on the available QuickBooks option.",
        "Confirm the action. Do not reuse the person's account or credentials.",
        "Review connected applications, banking access, payroll access, and other systems separately.",
        "Retain the completed removal checklist and updated user listing.",
    ])

    add_heading(doc, "Security safeguards and operating instructions", 1)

    safeguards = [
        ("1", "Enable two-step verification for every user", [
            "Have each user sign in to the Intuit Account manager using their own account.",
            "Open Sign in and security.",
            "Turn on two-step verification.",
            "Set up an authenticator application where available; otherwise use the approved phone or email method.",
            "Record the date completed without recording the user's secret or one-time code.",
            "Repeat for every active QuickBooks user and retain a signed or emailed confirmation.",
        ]),
        ("2", "Close the books after every monthly close", [
            "The Controller completes reconciliations and reviews the month-end financial statements.",
            "Open Settings, Account and settings, and Advanced.",
            "Open the Accounting section and enable Close the books.",
            "Set the approved closing date.",
            "Require a password or warning control for changes to closed periods, according to the available setting.",
            "Restrict knowledge of the closing password and document every approved override.",
            "Retain a screenshot of the closing date and the Controller's approval.",
        ]),
        ("3", "Review the Audit Log monthly", [
            "Open Settings and select Audit log.",
            "Filter the review period and review all users or the relevant administrator and accounting users.",
            "Review user and permission changes, settings changes, vendor changes, deleted or modified transactions, journal entries, reconciliation changes, and unusual administrator activity.",
            "Investigate unexpected entries and document the resolution.",
            "Have someone other than the primary transaction preparer sign off on the review.",
            "Save the available report, printout, PDF, or screenshots together with the review checklist.",
        ]),
        ("4", "Perform quarterly user access reviews", [
            "Export or capture the current QuickBooks user and role listing.",
            "Ask each department owner to confirm each user's continued need and correct role.",
            "Verify Primary Admin, Company Admin, payment, payroll, journal, and reporting access separately.",
            "Remove or reduce access that is no longer required.",
            "Document exceptions, owners, and remediation dates.",
            "Obtain management approval and retain the final before-and-after listing.",
        ]),
        ("5", "Control vendor and banking changes", [
            "Require a written vendor request supported by known vendor contact information.",
            "Independently verify banking changes using a previously known phone number or contact, not the information in the change request.",
            "Keep vendor maintenance separate from payment release whenever QuickBooks permissions allow.",
            "Require a second reviewer before paying a recently changed vendor.",
            "Review vendor changes through the Audit Log each month.",
            "Retain the request, verification evidence, approver, and first-payment review.",
        ]),
        ("6", "Control journal entries", [
            "Require support showing business purpose, accounts, amount, date, period, preparer, and proposed entry.",
            "Require independent approval for non-standard, correction, period-end, management, and material entries.",
            "Keep the preparer and approver separate whenever staffing permits.",
            "Review the Journal report and Audit Log during monthly close.",
            "Investigate unusual users, dates, accounts, round-dollar entries, late postings, and privileged-user entries.",
            "Retain at least one standard and one non-standard sample with approval evidence for SAX.",
        ]),
        ("7", "Control administrator access", [
            "Keep one Primary Admin and one named backup Company Admin.",
            "Assign all other users restricted roles.",
            "Do not use shared administrator accounts.",
            "Do not use administrator accounts for routine transaction processing.",
            "Review administrator activity monthly and administrator membership at least quarterly.",
            "Document any temporary administrator assignment and its expiration date.",
        ]),
    ]
    for number, title_text, steps in safeguards:
        add_heading(doc, f"Safeguard {number}  {title_text}", 2)
        add_numbered_steps(doc, steps)
        add_table(doc, ["Completion item", "Your response"], [
            ["Status", "[ ] Not started   [ ] In progress   [ ] Complete   [ ] Not applicable"],
            ["Owner", ""],
            ["Completion date", ""],
            ["Evidence file or location", ""],
            ["Open issue", ""],
        ], [2.1, 4.9], font_size=9.2, response_cols={1})

    add_heading(doc, "QuickBooks evidence package for SAX", 1)
    evidence_items = [
        ["1", "Current user and role listing", "Screenshot or available export from Manage users", ""],
        ["2", "Detailed custom-role permissions", "Permission summary for every role assigned", ""],
        ["3", "Role-change approvals", "Management approval and before-and-after access evidence", ""],
        ["4", "Two-step verification", "Confirmation from every active user", ""],
        ["5", "Primary and backup administrator approval", "Named approval with titles and date", ""],
        ["6", "QuickBooks Audit Log review", "Monthly review evidence and documented exceptions", ""],
        ["7", "Closed-period configuration", "Closing date screenshot and Controller approval", ""],
        ["8", "Quarterly access review", "Signed review with removed or changed access", ""],
        ["9", "Journal-entry samples", "One standard and one non-standard entry with support and approval", ""],
        ["10", "Vendor-change sample", "Request, independent verification, approval, and payment review", ""],
        ["11", "Bank reconciliation sample", "Completed reconciliation with independent reviewer sign-off", ""],
        ["12", "Role testing", "Evidence that allowed and prohibited actions were tested", ""],
    ]
    add_table(doc, ["No", "Evidence", "What to retain", "Complete and file name"], evidence_items, [0.4, 1.7, 3.45, 1.45], font_size=8.5, response_cols={3})

    add_heading(doc, "Accounting export already provided", 1)
    add_text(doc, "The ZIP dated 18 September 2026 contains the balance sheet, customer list, employee list, general ledger, accounting journal, profit and loss, vendor list, and trial balance. These support financial testing but do not identify who created, approved, changed, or deleted transactions.")
    add_text(doc, "The Journal spreadsheet is the accounting journal, not the QuickBooks security Audit Log. Collect the Audit Log separately using the instructions above.", bold_lead="The Journal spreadsheet is the accounting journal, not the QuickBooks security Audit Log. ")

    add_heading(doc, "Management sign off", 1)
    add_table(doc, ["Certification", "Name title date and response"], [
        ["Primary Admin confirmed", ""],
        ["Backup administrator confirmed", ""],
        ["Finance Controller approved role design", ""],
        ["Management approved user assignments", ""],
        ["Two-step verification confirmations collected", ""],
        ["Role testing completed", ""],
        ["Evidence package complete", ""],
        ["Open exceptions accepted or assigned", ""],
    ], [2.45, 4.55], font_size=9.3, response_cols={1})

    add_heading(doc, "Official QuickBooks references", 1)
    refs = [
        ("Add and manage custom roles", "https://quickbooks.intuit.com/learn-support/en-us/help-article/access-permissions/add-manage-custom-roles-quickbooks-online-advanced/L8Ugph7xl_US_en_US"),
        ("User roles and access rights", "https://quickbooks.intuit.com/learn-support/en-us/help-article/access-permissions/user-roles-access-rights-quickbooks-online/L66POfRrI_US_en_US"),
        ("Change the Primary Admin", "https://quickbooks.intuit.com/learn-support/en-us/help-article/primary-administrator/change-primary-admin-user-quickbooks-online/L9TU91iOk_US_en_US"),
        ("Use the QuickBooks Audit Log", "https://quickbooks.intuit.com/learn-support/en-us/help-article/audit-log/use-audit-log-quickbooks-online/L2WoVnW6I_US_en_US"),
        ("Set two-step verification", "https://quickbooks.intuit.com/learn-support/en-us/help-article/security-risk/verify-account-multi-factor-authentication/L2Xp0GNiT_US_en_US"),
        ("Set and protect a closing date", "https://quickbooks.intuit.com/learn-support/en-us/help-article/close-books/closing-dates-enter-dates-password-protecting/L5VWTMOBf_US_en_US"),
    ]
    for label, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_hyperlink(p, label, url)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("urbanXtracts  |  QuickBooks access safeguards  |  Management completion required")
    font_run(r, size=8.5, color=MID_GRAY)

    doc.core_properties.title = "QuickBooks Access Safeguards Implementation Workbook"
    doc.core_properties.subject = "QuickBooks role design, implementation instructions, and SAX audit evidence checklist"
    doc.core_properties.author = "urbanXtracts"
    doc.save(OUTFILE)
    print(OUTFILE)


if __name__ == "__main__":
    build()
